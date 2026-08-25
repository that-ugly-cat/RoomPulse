"""RoomPulse — FastAPI app.

Loop live: il presenter crea/usa un run, attiva una slide e ne cambia lo stato;
il pubblico entra col join_code (fisso sulla presentation → risolto al run attivo),
vede la slide attiva via polling, e vota. Aggregazione on-the-fly.

Dev run:  uv run uvicorn app.main:app --reload --port 8080
Seed:     uv run python seed.py
"""

import io
import json
import os
import re
import secrets
from contextlib import asynccontextmanager
from urllib.parse import urlparse

import segno
from openpyxl import Workbook
from fastapi import Cookie, Depends, FastAPI, HTTPException, Request, Response
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from pathlib import Path

from app import db, auth, locales, cluster as clustering
from app.aggregate import aggregate, SINGLE_VOTE_TYPES, MODERATED_TYPES
from app.mcp_app import mcp
from mcp.server.transport_security import TransportSecuritySettings

# dependency riusabile per le rotte presenter (alza 401 se non autenticato)
CurrentUser = Depends(auth.get_current_user)


def _require_admin(user: dict = CurrentUser) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(403, "richiede privilegi admin")
    return user


AdminUser = Depends(_require_admin)

STATIC_DIR = Path(__file__).resolve().parent / "static"

# ── Tier / API centralizzata ─────────────────────────────────────────────────
CENTRAL_API_KEY = os.environ.get("ANTHROPIC_API_KEY")   # chiave del server per i tier full/admin
# prezzi Sonnet (USD per 1M token), sovrascrivibili da env; solo per il tracking (niente blocco)
PRICE_IN = float(os.environ.get("RP_PRICE_IN_PER_MTOK", "3.0"))
PRICE_OUT = float(os.environ.get("RP_PRICE_OUT_PER_MTOK", "15.0"))
RP_ADMIN_EMAILS = [e.strip().lower() for e in os.environ.get("RP_ADMIN_EMAILS", "").split(",") if e.strip()]

# ── Moonshot (energizer collaborativo) ───────────────────────────────────────
MOON_DISTANCE_KM = 384400   # distanza media Terra-Luna, per l'altitudine
MOONSHOT_MIN_PLAYERS = 3

@asynccontextmanager
async def lifespan(app: FastAPI):
    db.init_db()
    if RP_ADMIN_EMAILS:  # promuove ad admin gli account elencati in env (bootstrap del primo admin)
        with db.get_conn() as conn:
            for em in RP_ADMIN_EMAILS:
                conn.execute("UPDATE user SET role='admin' WHERE email=?", (em,))
    async with mcp.session_manager.run():   # superficie MCP montata su /mcp
        yield


app = FastAPI(title="RoomPulse", lifespan=lifespan)

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


# ── Superficie MCP ───────────────────────────────────────────────────────────
# Il trasporto controlla l'header Host contro il DNS rebinding, quindi il dominio
# pubblico va dichiarato o ogni richiesta che passa da Caddy viene rifiutata.
def _allowed_hosts() -> list[str]:
    # `:*` e' un carattere jolly sulla porta, non sull'host: in sviluppo la porta
    # cambia a ogni avvio, mentre il dominio pubblico resta uno solo ed esatto.
    hosts = ["localhost", "127.0.0.1", "localhost:*", "127.0.0.1:*"]
    public = urlparse(os.environ.get("PUBLIC_URL", "")).netloc
    if public:
        hosts.append(public)
    return hosts


app.mount("/mcp", mcp.streamable_http_app(
    streamable_http_path="/", json_response=True, stateless_http=True,
    transport_security=TransportSecuritySettings(
        allowed_hosts=_allowed_hosts(),
        allowed_origins=[os.environ.get("PUBLIC_URL", ""), "http://localhost:*",
                         "http://127.0.0.1:*"])))


@app.middleware("http")
async def mcp_key_gate(request: Request, call_next):
    """Risolve il chiamante MCP, o rifiuta.

    Due strade, una tabella: l'header e' la via normale, `/mcp/k/{key}` porta la
    stessa chiave come segmento di path per i client che non sanno mandare header,
    e viene tolto prima che l'app montata lo veda — cosi' il layer MCP non sa
    nemmeno come si e' autenticato chi lo chiama.

    Nota: questa e' l'UNICA porta di `/mcp`. Il gate SSO davanti all'app non
    interviene qui (un client MCP non ha un browser con cui fare login), quindi
    il blocco Caddy deve lasciar passare `/mcp` senza forward_auth."""
    path = request.url.path
    if not path.startswith("/mcp"):
        return await call_next(request)

    if path.startswith("/mcp/k/"):
        key, _, rest = path[len("/mcp/k/"):].partition("/")
        request.scope["path"] = "/mcp/" + rest
        request.scope["raw_path"] = request.scope["path"].encode()
    else:
        key = request.headers.get("X-API-Key", "")

    caller = auth.check_mcp_key(key)
    auth.set_mcp_caller(caller)
    if not caller:
        return JSONResponse({"error": "chiave API mancante o non valida"}, status_code=401)
    return await call_next(request)


# ----------------------------------------------------------------------------
# Pagine
# ----------------------------------------------------------------------------
@app.get("/")
def audience_page():
    return FileResponse(STATIC_DIR / "audience.html")


def _al_login():
    """Dove mandare chi non e' autenticato su una pagina che lo richiede.

    In `gateway` **non** si rimanda a `/login`: quella rotta, in questa
    modalita', l'app la spegne da se' e rimanda a `/edit` — i due si
    rimbalzerebbero all'infinito. In produzione non capita, perche' il gate
    intercetta prima che la richiesta arrivi qui; ma se il matcher del proxy
    fosse sbagliato si girerebbe a vuoto invece di ricevere un errore, e un
    anello e' molto piu' difficile da diagnosticare di un codice di stato.

    Il messaggio e' quello di Onopedia: parla all'**operatore**, perche' in
    `gateway` una richiesta senza identita' significa che il gate non ha
    girato — un guasto di configurazione, non della persona.
    """
    if auth.gateway_mode():
        raise HTTPException(status_code=503, detail=(
            "Gateway mode: no valid identity in the X-Borant-* headers. Check "
            "that the gate really sits in front of this app and that "
            "BORANT_TRUSTED_PROXY lists the address the proxy connects from."))
    return RedirectResponse("/login")


@app.get("/present")
def presenter_page(request: Request, session: str | None = Cookie(default=None)):
    if not auth.get_user_or_none(session, request):
        return _al_login()
    return FileResponse(STATIC_DIR / "presenter.html")


@app.get("/edit")
def editor_page(request: Request, session: str | None = Cookie(default=None)):
    if not auth.get_user_or_none(session, request):
        return _al_login()
    return FileResponse(STATIC_DIR / "editor.html")


@app.get("/admin")
def admin_page(request: Request, session: str | None = Cookie(default=None)):
    u = auth.get_user_or_none(session, request)
    if not u:
        return _al_login()
    if u.get("role") != "admin":
        return RedirectResponse("/edit")
    return FileResponse(STATIC_DIR / "admin.html")


@app.get("/login")
def login_page(request: Request, session: str | None = Cookie(default=None)):
    # In `gateway` l'app spegne il proprio login da sola invece di affidarsi al
    # proxy per nasconderlo: conosce la propria modalita' meglio del reverse
    # proxy, e due anagrafiche in parallelo vorrebbero dire SSO non imposto.
    #
    # La destinazione e' /edit e non "/": "/" e' la pagina del PUBBLICO, e chi
    # arriva su /login e' un presenter. Mandarlo alla schermata d'ingresso del
    # pubblico sarebbe atterrare nel posto sbagliato dopo aver fatto il login.
    if auth.gateway_mode():
        return RedirectResponse("/edit")
    return FileResponse(STATIC_DIR / "login.html")


@app.get("/guide")
def guide_page():
    return FileResponse(STATIC_DIR / "guide.html")


class LoginIn(BaseModel):
    email: str
    password: str


SIGNUP_CODE = os.environ.get("RP_SIGNUP_CODE")  # se settato, la registrazione lo richiede


@app.get("/api/auth-config")
def auth_config():
    return {"signup_code_required": bool(SIGNUP_CODE)}


MIN_PASSWORD = 6


def _clean_email(raw: str) -> str:
    email = (raw or "").strip().lower()
    if not email or "@" not in email:
        raise HTTPException(400, "email non valida")
    return email


def _check_password(pw: str) -> str:
    if len(pw or "") < MIN_PASSWORD:
        raise HTTPException(400, f"password troppo corta (minimo {MIN_PASSWORD} caratteri)")
    return pw


def _bump_token_version(conn, uid: str) -> int:
    """Invalida le sessioni gia aperte di quell'utente. Senza questo un reset password
    sarebbe cosmetico: il cookie vecchio resterebbe valido per giorni."""
    conn.execute("UPDATE user SET token_version = token_version + 1 WHERE id=?", (uid,))
    return conn.execute("SELECT token_version FROM user WHERE id=?", (uid,)).fetchone()[0]


class RegisterIn(BaseModel):
    email: str
    password: str
    name: str = ""
    signup_code: str | None = None


@app.post("/api/register")
def register(body: RegisterIn, response: Response):
    # In `gateway` gli account nascono nel gate, non qui: lasciare aperta questa
    # rotta vorrebbe dire due anagrafiche in parallelo e l'SSO non imposto.
    if auth.gateway_mode():
        raise HTTPException(403, "Registrazione gestita dal gate SSO")
    if SIGNUP_CODE and (body.signup_code or "") != SIGNUP_CODE:
        raise HTTPException(403, "codice di registrazione non valido")
    email = _clean_email(body.email)
    _check_password(body.password)
    with db.get_conn() as conn:
        if conn.execute("SELECT 1 FROM user WHERE email=?", (email,)).fetchone():
            raise HTTPException(409, "email già registrata")
        uid = db.new_id()
        role = "admin" if email in RP_ADMIN_EMAILS else "free"
        conn.execute(
            "INSERT INTO user (id, email, password_hash, name, is_active, role, created_at) "
            "VALUES (?,?,?,?,1,?,?)",
            (uid, email, auth.hash_password(body.password),
             body.name.strip() or email.split("@")[0], role, db.now_iso()),
        )
    token = auth.create_token(uid)
    response.set_cookie(
        "session", token, httponly=True, samesite="lax", max_age=auth.EXPIRE_DAYS * 86400
    )
    return {"ok": True}


@app.post("/api/login")
def login(body: LoginIn, response: Response):
    if auth.gateway_mode():
        raise HTTPException(403, "Accesso gestito dal gate SSO")
    with db.get_conn() as conn:
        u = conn.execute(
            "SELECT * FROM user WHERE email=? AND is_active=1", (body.email,)
        ).fetchone()
    if not u or not auth.verify_password(body.password, u["password_hash"]):
        raise HTTPException(401, "Credenziali errate")
    token = auth.create_token(u["id"], u["token_version"])
    response.set_cookie(
        "session", token, httponly=True, samesite="lax", max_age=auth.EXPIRE_DAYS * 86400
    )
    return {"ok": True, "name": u["name"]}


BORANT_LOGOUT_URL = os.environ.get("BORANT_LOGOUT_URL", "https://id.borant.eu/logout")


@app.post("/api/logout")
def logout(response: Response):
    response.delete_cookie("session")
    # In `gateway` buttare il cookie locale non e' uscire: la sessione sta nel
    # gate e il click successivo rientra da solo. Il client manda il browser li'.
    if auth.gateway_mode():
        return {"ok": True, "redirect": BORANT_LOGOUT_URL}
    return {"ok": True}


@app.get("/api/me")
def me(user: dict = CurrentUser):
    with db.get_conn() as conn:
        row = conn.execute("SELECT api_key FROM user WHERE id=?", (user["id"],)).fetchone()
    key = row["api_key"] if row else None
    # non rivelo la chiave: solo se c'è e gli ultimi 4 char
    masked = ("…" + key[-4:]) if key else None
    role = user.get("role", "free")
    return {"email": user["email"], "name": user["name"], "role": role,
            "uses_central": role in ("full", "admin"), "central_available": bool(CENTRAL_API_KEY),
            "api_key_set": bool(key), "api_key_hint": masked}


class ApiKeyIn(BaseModel):
    api_key: str


@app.put("/api/me/api-key")
def set_api_key(body: ApiKeyIn, user: dict = CurrentUser):
    key = body.api_key.strip()
    with db.get_conn() as conn:
        conn.execute("UPDATE user SET api_key=? WHERE id=?", (key or None, user["id"]))
    return {"ok": True, "api_key_set": bool(key)}


# ── Chiavi della superficie MCP ──────────────────────────────────────────────
# Sono chiavi *di una persona*: ogni chiamata MCP gira come il proprietario e
# raggiunge esattamente le sue deck. Da non confondere con `user.api_key`, che e'
# la chiave Anthropic per il clustering — altro scopo, altra tabella.
class McpKeyIn(BaseModel):
    name: str = "mcp"


@app.get("/api/me/mcp-keys")
def list_mcp_keys(user: dict = CurrentUser):
    """Le chiavi del chiamante. Il valore NON si rilegge: mostrato una volta sola
    alla creazione, qui restano nome, prefisso e ultimo uso."""
    with db.get_conn() as conn:
        rows = conn.execute(
            "SELECT id, name, key, active, created_at, last_used_at FROM mcp_key "
            "WHERE user_id=? ORDER BY created_at DESC", (user["id"],),
        ).fetchall()
    return [{"id": r["id"], "name": r["name"], "hint": r["key"][:6] + "…" + r["key"][-4:],
             "active": bool(r["active"]), "created_at": r["created_at"],
             "last_used_at": r["last_used_at"]} for r in rows]


@app.post("/api/me/mcp-keys")
def create_mcp_key(body: McpKeyIn, user: dict = CurrentUser):
    """Crea una chiave e la restituisce in chiaro **una volta sola**."""
    key = auth.new_mcp_key()
    with db.get_conn() as conn:
        conn.execute(
            "INSERT INTO mcp_key (id, user_id, name, key, active, created_at) VALUES (?,?,?,?,1,?)",
            (db.new_id(), user["id"], body.name.strip() or "mcp", key, db.now_iso()),
        )
    return {"key": key, "name": body.name.strip() or "mcp"}


@app.delete("/api/me/mcp-keys/{kid}")
def revoke_mcp_key(kid: str, user: dict = CurrentUser):
    """Revoca: la riga resta (per `last_used_at`), la chiave smette di aprire."""
    with db.get_conn() as conn:
        row = conn.execute("SELECT user_id FROM mcp_key WHERE id=?", (kid,)).fetchone()
        if not row or row["user_id"] != user["id"]:
            raise HTTPException(404, "chiave non trovata")
        conn.execute("UPDATE mcp_key SET active=0 WHERE id=?", (kid,))
    return {"ok": True}


class PasswordChangeIn(BaseModel):
    current: str
    new: str


@app.put("/api/me/password")
def change_own_password(body: PasswordChangeIn, response: Response, user: dict = CurrentUser):
    """Cambio password proprio. Richiede quella attuale, invalida le altre sessioni
    e ri-emette il cookie di QUESTA, altrimenti chi cambia password si sloggherebbe da solo."""
    _check_password(body.new)
    with db.get_conn() as conn:
        row = conn.execute("SELECT password_hash FROM user WHERE id=?", (user["id"],)).fetchone()
        if not row or not auth.verify_password(body.current, row["password_hash"]):
            raise HTTPException(403, "password attuale errata")
        conn.execute("UPDATE user SET password_hash=? WHERE id=?",
                     (auth.hash_password(body.new), user["id"]))
        ver = _bump_token_version(conn, user["id"])
    response.set_cookie("session", auth.create_token(user["id"], ver),
                        httponly=True, samesite="lax", max_age=auth.EXPIRE_DAYS * 86400)
    return {"ok": True}


# ── Admin: gestione utenti + costi ───────────────────────────────────────────
@app.get("/api/admin/users")
def admin_users(user: dict = AdminUser):
    with db.get_conn() as conn:
        rows = conn.execute(
            "SELECT u.id, u.email, u.name, u.role, u.is_active, u.created_at, "
            "  (u.api_key IS NOT NULL) AS own_key, "
            "  (SELECT COUNT(*) FROM presentation p WHERE p.owner=u.id) AS n_decks, "
            "  COALESCE((SELECT SUM(cost_usd) FROM usage_log WHERE user_id=u.id),0) AS cost, "
            "  COALESCE((SELECT SUM(input_tokens+output_tokens) FROM usage_log WHERE user_id=u.id),0) AS tokens "
            "FROM user u ORDER BY u.created_at DESC"
        ).fetchall()
        tot = conn.execute(
            "SELECT COALESCE(SUM(cost_usd),0) AS c, COALESCE(SUM(input_tokens+output_tokens),0) AS t FROM usage_log"
        ).fetchone()
    return {"users": [dict(r) for r in rows], "total_cost": tot["c"], "total_tokens": tot["t"],
            "central_available": bool(CENTRAL_API_KEY)}


class AdminUserPatch(BaseModel):
    role: str | None = None
    is_active: bool | None = None


@app.patch("/api/admin/users/{uid}")
def admin_set_user(uid: str, body: AdminUserPatch, user: dict = AdminUser):
    with db.get_conn() as conn:
        target = conn.execute("SELECT id, role FROM user WHERE id=?", (uid,)).fetchone()
        if not target:
            raise HTTPException(404, "utente non trovato")

        def _last_admin_guard():
            n = conn.execute(
                "SELECT COUNT(*) AS c FROM user WHERE role='admin' AND is_active=1"
            ).fetchone()["c"]
            if n <= 1:
                raise HTTPException(400, "non puoi rimuovere/disattivare l'ultimo admin")

        if body.role is not None:
            if body.role not in ("free", "full", "admin"):
                raise HTTPException(400, "ruolo non valido")
            if target["role"] == "admin" and body.role != "admin":
                _last_admin_guard()
            conn.execute("UPDATE user SET role=? WHERE id=?", (body.role, uid))
        if body.is_active is not None:
            if target["role"] == "admin" and not body.is_active:
                _last_admin_guard()
            conn.execute("UPDATE user SET is_active=? WHERE id=?", (1 if body.is_active else 0, uid))
        return {"ok": True}


class AdminUserCreate(BaseModel):
    email: str
    password: str
    name: str = ""
    role: str = "free"


@app.post("/api/admin/users")
def admin_create_user(body: AdminUserCreate, user: dict = AdminUser):
    """Crea un utente dal pannello, senza passare dall'autoregistrazione o dalla shell."""
    email = _clean_email(body.email)
    _check_password(body.password)
    if body.role not in ("free", "full", "admin"):
        raise HTTPException(400, "ruolo non valido")
    with db.get_conn() as conn:
        if conn.execute("SELECT 1 FROM user WHERE email=?", (email,)).fetchone():
            raise HTTPException(409, "email gia registrata")
        uid = db.new_id()
        conn.execute(
            "INSERT INTO user (id, email, password_hash, name, is_active, role, created_at) "
            "VALUES (?,?,?,?,1,?,?)",
            (uid, email, auth.hash_password(body.password),
             body.name.strip() or email.split("@")[0], body.role, db.now_iso()),
        )
    return {"ok": True, "id": uid, "email": email}


class AdminPasswordIn(BaseModel):
    password: str


@app.put("/api/admin/users/{uid}/password")
def admin_reset_password(uid: str, body: AdminPasswordIn, user: dict = AdminUser):
    """Reimposta la password di un altro utente e ne butta giu le sessioni aperte."""
    _check_password(body.password)
    with db.get_conn() as conn:
        if not conn.execute("SELECT 1 FROM user WHERE id=?", (uid,)).fetchone():
            raise HTTPException(404, "utente non trovato")
        conn.execute("UPDATE user SET password_hash=? WHERE id=?",
                     (auth.hash_password(body.password), uid))
        _bump_token_version(conn, uid)
    return {"ok": True}


@app.delete("/api/admin/users/{uid}")
def admin_delete_user(uid: str, user: dict = AdminUser):
    """Cancella un utente e TUTTE le sue deck (slide, run, risposte, voti, cluster).
    Le righe di usage_log restano: servono a tenere corretto il totale dei costi."""
    if uid == user["id"]:
        raise HTTPException(400, "non puoi cancellare te stesso")
    with db.get_conn() as conn:
        target = conn.execute("SELECT id, role FROM user WHERE id=?", (uid,)).fetchone()
        if not target:
            raise HTTPException(404, "utente non trovato")
        if target["role"] == "admin":
            # conta gli ALTRI admin attivi: includere il bersaglio bloccherebbe anche
            # la cancellazione di un admin gia disattivato
            n = conn.execute(
                "SELECT COUNT(*) AS c FROM user WHERE role='admin' AND is_active=1 AND id<>?",
                (uid,),
            ).fetchone()["c"]
            if n < 1:
                raise HTTPException(400, "non puoi cancellare l'ultimo admin")
        pids = [r["id"] for r in conn.execute(
            "SELECT id FROM presentation WHERE owner=?", (uid,)).fetchall()]
        for pid in pids:
            _purge_presentation(conn, pid)
        conn.execute("DELETE FROM user WHERE id=?", (uid,))
        return {"ok": True, "deleted_decks": len(pids)}


def _remap_carry(conn, new_slide_id: str, config_json: str, idmap: dict) -> None:
    """argstep: `carry_from` e un id di slide, quindi dopo un clone o un import va
    rimappato ai nuovi id esattamente come `pair_id`. Se la sorgente non e nella deck
    la catena resta scollegata (campi editabili a runtime) invece di puntare nel vuoto."""
    cfg = json.loads(config_json)
    if not cfg.get("carry_from"):
        return
    cfg["carry_from"] = idmap.get(cfg["carry_from"])
    conn.execute("UPDATE slide SET config=? WHERE id=?", (json.dumps(cfg), new_slide_id))


def _clone_presentation(conn, source_pid: str, owner: str) -> str:
    """Duplica una deck (slide + config + note + pair, niente run) in un nuovo account."""
    src = conn.execute("SELECT * FROM presentation WHERE id=?", (source_pid,)).fetchone()
    slides = conn.execute(
        "SELECT * FROM slide WHERE presentation_id=? ORDER BY ord", (source_pid,)
    ).fetchall()
    new_pid = db.new_id()
    conn.execute(
        "INSERT INTO presentation (id, title, owner, join_code, created_at) VALUES (?,?,?,?,?)",
        (new_pid, src["title"], owner, db.new_join_code(conn), db.now_iso()),
    )
    idmap: dict = {}
    for s in slides:
        nid = db.new_id()
        conn.execute(
            "INSERT INTO slide (id, presentation_id, ord, type, question, config, presenter_notes) "
            "VALUES (?,?,?,?,?,?,?)",
            (nid, new_pid, s["ord"], s["type"], s["question"], s["config"], s["presenter_notes"]),
        )
        idmap[s["id"]] = nid
    for s in slides:  # remap pre/post e catena argstep ai nuovi id
        if s["pair_id"] and s["pair_id"] in idmap:
            conn.execute("UPDATE slide SET pair_id=? WHERE id=?", (idmap[s["pair_id"]], idmap[s["id"]]))
        if s["type"] == "argstep":
            _remap_carry(conn, idmap[s["id"]], s["config"], idmap)
    return new_pid


class DistributeIn(BaseModel):
    source_pid: str
    target_user_ids: list[str]


@app.post("/api/admin/distribute")
def admin_distribute(body: DistributeIn, user: dict = AdminUser):
    """Clona una deck dell'admin negli account degli utenti scelti (per chi non sa importare un JSON)."""
    with db.get_conn() as conn:
        src = conn.execute("SELECT owner FROM presentation WHERE id=?", (body.source_pid,)).fetchone()
        if not src:
            raise HTTPException(404, "deck sorgente non trovata")
        if src["owner"] != user["id"]:
            raise HTTPException(403, "puoi distribuire solo deck che possiedi")
        n = 0
        for uid in body.target_user_ids:
            if conn.execute("SELECT 1 FROM user WHERE id=? AND is_active=1", (uid,)).fetchone():
                _clone_presentation(conn, body.source_pid, uid)
                n += 1
        return {"ok": True, "distributed_to": n}


@app.get("/api/i18n")
def i18n(lang: str = locales.DEFAULT):
    """Stringhe UI per la lingua richiesta (aperto, serve anche all'audience)."""
    chosen = lang if lang in locales.SUPPORTED else locales.DEFAULT
    return {"lang": chosen, "supported": list(locales.SUPPORTED), "t": locales.get_t(chosen)}


@app.get("/qr/{code}.svg")
def qr_svg(code: str, request: Request):
    base = str(request.base_url).rstrip("/")
    url = f"{base}/?c={code}"
    buff = io.BytesIO()
    segno.make(url).save(buff, kind="svg", scale=6, border=2)
    return Response(content=buff.getvalue(), media_type="image/svg+xml")


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def _slide_dict(row) -> dict:
    return {
        "id": row["id"],
        "ord": row["ord"],
        "type": row["type"],
        "question": row["question"],
        "config": json.loads(row["config"]),
        "pair_id": row["pair_id"],
        "presenter_notes": row["presenter_notes"],
    }


def _run_slide_state(conn, run_id: str, slide_id: str) -> str:
    row = conn.execute(
        "SELECT state FROM run_slide WHERE run_id=? AND slide_id=?",
        (run_id, slide_id),
    ).fetchone()
    return row["state"] if row else "pending"


# --- headcount: chi c'e' in sala, e chi ha davvero votato ---------------------
# La presenza e' un heartbeat scritto dal poll del pubblico (/api/live), non dal voto:
# conta le persone COLLEGATE, comprese quelle che stanno zitte.
PRESENCE_TTL_MS = 20_000    # oltre questa soglia il token e' considerato uscito
PRESENCE_WRITE_MS = 5_000   # il poll gira ogni 1.5s: non riscrivo a ogni giro


def _touch_presence(conn, run_id: str, token: str) -> None:
    now = db.now_ms()
    conn.execute(
        "INSERT INTO presence (run_id, token, last_seen_ms) VALUES (?,?,?) "
        "ON CONFLICT(run_id, token) DO UPDATE SET last_seen_ms=excluded.last_seen_ms "
        "WHERE excluded.last_seen_ms - presence.last_seen_ms > ?",
        (run_id, token, now, PRESENCE_WRITE_MS),
    )


def _present_count(conn, run_id: str) -> int:
    return conn.execute(
        "SELECT COUNT(*) AS c FROM presence WHERE run_id=? AND last_seen_ms >= ?",
        (run_id, db.now_ms() - PRESENCE_TTL_MS),
    ).fetchone()["c"]


def _voter_count(conn, run_id: str, slide_id: str) -> int:
    """PERSONE distinte che hanno risposto, non righe di risposta.

    La differenza conta in due casi opposti, ed e' il motivo per cui questo numero
    non e' `results.n`:
    - risposta multipla (mc con `multi`): una persona, una riga, ma N opzioni segnate,
      quindi le barre sommano piu' dei votanti;
    - tipi a invii ripetuti (opentext, wordcloud, qa, argpoll senza `single`): una
      persona, N righe, quindi le risposte sono piu' dei votanti.
    Le risposte nascoste dalla moderazione restano contate: la persona ha votato."""
    return conn.execute(
        "SELECT COUNT(DISTINCT participant_token) AS c FROM response "
        "WHERE run_id=? AND slide_id=?",
        (run_id, slide_id),
    ).fetchone()["c"]


def _merged_mc_options(conn, run_id: str, slide) -> list:
    """Opzioni base (config) + quelle aggiunte dai partecipanti in questo run."""
    config = json.loads(slide["config"])
    options = [dict(o) for o in config.get("options", [])]
    extra = conn.execute(
        "SELECT id, label FROM mc_option WHERE run_id=? AND slide_id=? ORDER BY created_at",
        (run_id, slide["id"]),
    ).fetchall()
    options += [{"id": r["id"], "label": r["label"], "added": True} for r in extra]
    return options


def _mc_results(conn, run_id: str, slide) -> dict:
    """mc: conta choice (single) e choices[] (multiple), su opzioni base + aggiunte."""
    config = json.loads(slide["config"])
    options = _merged_mc_options(conn, run_id, slide)
    rows = conn.execute(
        "SELECT payload FROM response WHERE run_id=? AND slide_id=? AND status='visible'",
        (run_id, slide["id"]),
    ).fetchall()
    counts: dict = {}
    n = 0
    for r in rows:
        p = json.loads(r["payload"])
        n += 1
        picks = p["choices"] if isinstance(p.get("choices"), list) else (
            [p["choice"]] if p.get("choice") is not None else []
        )
        for c in picks:
            counts[c] = counts.get(c, 0) + 1
    return {
        "type": "mc",
        "n": n,
        "multi": bool(config.get("multi")),
        "quiz": bool(config.get("quiz")),
        "correct": config.get("correct", []),
        "options": [
            {"id": o["id"], "label": o["label"], "count": counts.get(o["id"], 0),
             "added": o.get("added", False)}
            for o in options
        ],
    }


def _results(conn, run_id: str, slide) -> dict:
    if slide["type"] == "argstep":
        return _argstep_results(conn, run_id, slide)
    if slide["type"] == "qa":
        return _qa_results(conn, run_id, slide["id"])
    if slide["type"] == "mc":
        return _mc_results(conn, run_id, slide)
    if slide["type"] == "argpoll":
        has = conn.execute(
            "SELECT 1 FROM cluster WHERE run_id=? AND slide_id=? LIMIT 1",
            (run_id, slide["id"]),
        ).fetchone()
        if has:
            return _argpoll_clustered(conn, run_id, slide["id"])
    if slide["type"] == "opentext":
        has = conn.execute(
            "SELECT 1 FROM cluster WHERE run_id=? AND slide_id=? AND kind='theme' LIMIT 1",
            (run_id, slide["id"]),
        ).fetchone()
        if has:
            return _opentext_clustered(conn, run_id, slide["id"])
    rows = conn.execute(
        "SELECT id, payload, status, created_at FROM response "
        "WHERE run_id=? AND slide_id=? AND status='visible' ORDER BY created_at",
        (run_id, slide["id"]),
    ).fetchall()
    return aggregate(slide["type"], json.loads(slide["config"]), rows)


# --- argstep: la catena claim -> justification -> objection -------------------
# Tre slide distinte, non tre fasi di una sola: cosi ogni tappa ha il suo stato
# open/closed/revealed, la sua coda di moderazione e il suo momento di reveal.
# Il payload della tappa N porta con se una COPIA dei campi delle tappe precedenti,
# quindi la tappa finale contiene da sola l'argomento intero: si clusterizza li, una volta.
ARGSTEP_CHAIN = ("claim", "justification", "objection")


def _argstep_fields(cfg: dict) -> list:
    """La catena e canonica e ordinata: di `fields` conta solo la LUNGHEZZA.
    Cosi una config malformata degrada a un prefisso valido invece di rompere."""
    n = max(1, min(len(ARGSTEP_CHAIN), len(cfg.get("fields") or [])))
    return list(ARGSTEP_CHAIN[:n])


def _peer_source(conn, run_id: str, slide_id: str, src_id: str, token: str):
    """Assegna a `token` la risposta di QUALCUN ALTRO sulla slide sorgente, e la memorizza:
    deve restare stabile fra un poll e l'altro. Bilanciata: vince quella che ha
    ricevuto meno obiezioni finora."""
    ex = conn.execute(
        "SELECT source_response_id FROM carry_assign WHERE run_id=? AND slide_id=? AND token=?",
        (run_id, slide_id, token),
    ).fetchone()
    if ex:
        row = conn.execute(
            "SELECT id, payload FROM response WHERE id=? AND status='visible'",
            (ex["source_response_id"],),
        ).fetchone()
        if row:
            return row
        # sorgente nascosta dalla moderazione, o sostituita dall'upsert del voto singolo
        conn.execute(
            "DELETE FROM carry_assign WHERE run_id=? AND slide_id=? AND token=?",
            (run_id, slide_id, token),
        )
    pool = conn.execute(
        "SELECT id, payload FROM response WHERE run_id=? AND slide_id=? AND status='visible' "
        "AND participant_token<>? ORDER BY created_at",
        (run_id, src_id, token),
    ).fetchall()
    if not pool:
        return None
    load = {
        r["source_response_id"]: r["n"]
        for r in conn.execute(
            "SELECT source_response_id, COUNT(*) AS n FROM carry_assign "
            "WHERE run_id=? AND slide_id=? GROUP BY source_response_id",
            (run_id, slide_id),
        ).fetchall()
    }
    chosen = min(pool, key=lambda r: (load.get(r["id"], 0), r["id"]))  # tie-break stabile
    conn.execute(
        "INSERT OR REPLACE INTO carry_assign "
        "(run_id, slide_id, token, source_response_id, created_at) VALUES (?,?,?,?,?)",
        (run_id, slide_id, token, chosen["id"], db.now_iso()),
    )
    return chosen


def _resolve_carry(conn, run_id: str, slide, token: str | None) -> dict | None:
    """Cosa questa slide eredita dalla precedente, per questo token. None se non c'e ancora."""
    cfg = json.loads(slide["config"])
    src_id = cfg.get("carry_from")
    carried = _argstep_fields(cfg)[:-1]
    if not src_id or not carried or not token:
        return None
    mode = "peer" if cfg.get("carry_mode") == "peer" else "self"
    if mode == "peer":
        row = _peer_source(conn, run_id, slide["id"], src_id, token)
    else:
        row = conn.execute(
            "SELECT id, payload FROM response WHERE run_id=? AND slide_id=? AND participant_token=? "
            "AND status='visible' ORDER BY created_at DESC LIMIT 1",
            (run_id, src_id, token),
        ).fetchone()
    if not row:
        return None
    p = json.loads(row["payload"])
    return {"values": {f: p.get(f, "") for f in carried}, "source_id": row["id"], "mode": mode}


def _argstep_results(conn, run_id: str, slide) -> dict:
    """Lo stato finale che serve in aula: UNA RIGA PER ARGOMENTO, con le obiezioni
    (al plurale) raccolte sotto. Il raggruppamento e per risposta ereditata, quindi in
    modo peer piu persone che obiettano allo stesso claim finiscono nella stessa riga."""
    cfg = json.loads(slide["config"])
    fields = _argstep_fields(cfg)
    collect = fields[-1]
    sid = slide["id"]
    clabels = {
        c["id"]: c["label"]
        for c in conn.execute(
            "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=?", (run_id, sid)
        ).fetchall()
    }
    claim_cl = conn.execute(
        "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=? AND kind='claim' ORDER BY ord",
        (run_id, sid),
    ).fetchall()
    arg_cl = conn.execute(
        "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=? AND kind='arg' ORDER BY ord",
        (run_id, sid),
    ).fetchall()
    rows = conn.execute(
        "SELECT id, payload, source_response_id, claim_cluster_id, arg_cluster_id "
        "FROM response WHERE run_id=? AND slide_id=? AND status='visible' ORDER BY created_at",
        (run_id, sid),
    ).fetchall()

    out: list = []
    index: dict = {}
    matrix: dict = {(c["id"], a["id"]): 0 for c in claim_cl for a in arg_cl}
    for r in rows:
        p = json.loads(r["payload"])
        # chiave di raggruppamento: la risposta ereditata; senza carry, la risposta stessa
        key = r["source_response_id"] or r["id"]
        row = index.get(key)
        if row is None:
            row = {
                "key": key,
                "claim": p.get("claim", ""),
                "justification": p.get("justification", ""),
                "cluster": "",
                "cluster_id": None,
                "objections": [],
            }
            index[key] = row
            out.append(row)
        if not row["cluster_id"] and r["claim_cluster_id"]:
            # il claim e lo stesso per tutta la riga: basta la prima assegnazione non vuota
            row["cluster_id"] = r["claim_cluster_id"]
            row["cluster"] = clabels.get(r["claim_cluster_id"], "")
        if collect == "objection":
            row["objections"].append({
                "id": r["id"],
                "text": p.get("objection", ""),
                "by": p.get("by") or "",
                "tag": clabels.get(r["arg_cluster_id"], ""),
            })
        if (r["claim_cluster_id"], r["arg_cluster_id"]) in matrix:
            matrix[(r["claim_cluster_id"], r["arg_cluster_id"])] += 1

    order = {c["id"]: i for i, c in enumerate(claim_cl)}
    out.sort(key=lambda x: (order.get(x["cluster_id"], 10**6), -len(x["objections"])))
    res = {
        "type": "argstep",
        "n": len(rows),
        "fields": fields,
        "collect": collect,
        "labels": cfg.get("labels") or {},
        "carry_mode": "peer" if cfg.get("carry_mode") == "peer" else "self",
        "rows": out,
        "clustered": bool(claim_cl),
        "clusters": [
            {"id": c["id"], "label": c["label"],
             "count": sum(1 for r in out if r["cluster_id"] == c["id"])}
            for c in claim_cl
        ],
    }
    if claim_cl and arg_cl:  # matrice claim x tipo di obiezione, stessa forma di argpoll
        res["matrix"] = [[matrix[(c["id"], a["id"])] for a in arg_cl] for c in claim_cl]
        res["matrix_rows"] = [c["label"] for c in claim_cl]
        res["matrix_cols"] = [a["label"] for a in arg_cl]
    return res


def _argpoll_clustered(conn, run_id: str, slide_id: str) -> dict:
    """Vista clusterizzata: claim cluster annidati con tag dell'argomento + matrice claim×arg."""
    claim_cl = conn.execute(
        "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=? AND kind='claim' ORDER BY ord",
        (run_id, slide_id),
    ).fetchall()
    arg_cl = conn.execute(
        "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=? AND kind='arg' ORDER BY ord",
        (run_id, slide_id),
    ).fetchall()
    arg_label = {a["id"]: a["label"] for a in arg_cl}
    rows = conn.execute(
        "SELECT payload, claim_cluster_id, arg_cluster_id FROM response "
        "WHERE run_id=? AND slide_id=? AND status='visible'",
        (run_id, slide_id),
    ).fetchall()
    by_claim: dict = {c["id"]: [] for c in claim_cl}
    matrix: dict = {(c["id"], a["id"]): 0 for c in claim_cl for a in arg_cl}
    n = 0
    for r in rows:
        p = json.loads(r["payload"])
        n += 1
        cc, ac = r["claim_cluster_id"], r["arg_cluster_id"]
        if cc in by_claim:
            by_claim[cc].append({
                "claim": p.get("claim", ""),
                "justification": p.get("justification", ""),
                "arg_label": arg_label.get(ac, ""),
            })
        if (cc, ac) in matrix:
            matrix[(cc, ac)] += 1
    claims_out = sorted(
        [{"id": c["id"], "label": c["label"], "count": len(by_claim[c["id"]]),
          "items": by_claim[c["id"]]} for c in claim_cl],
        key=lambda x: -x["count"],
    )
    return {
        "type": "argpoll",
        "n": n,
        "clustered": True,
        "claim_clusters": claims_out,
        "arg_clusters": [{"id": a["id"], "label": a["label"]} for a in arg_cl],
        "matrix": [[matrix[(c["id"], a["id"])] for a in arg_cl] for c in claim_cl],
        "matrix_rows": [c["label"] for c in claim_cl],
        "matrix_cols": [a["label"] for a in arg_cl],
    }


def _opentext_clustered(conn, run_id: str, slide_id: str) -> dict:
    """Open text clusterizzato: cluster tematici (un asse) con le risposte annidate."""
    clusters = conn.execute(
        "SELECT id, label FROM cluster WHERE run_id=? AND slide_id=? AND kind='theme' ORDER BY ord",
        (run_id, slide_id),
    ).fetchall()
    rows = conn.execute(
        "SELECT payload, cluster_id FROM response "
        "WHERE run_id=? AND slide_id=? AND status='visible'",
        (run_id, slide_id),
    ).fetchall()
    by_cluster: dict = {c["id"]: [] for c in clusters}
    n = 0
    for r in rows:
        n += 1
        if r["cluster_id"] in by_cluster:
            by_cluster[r["cluster_id"]].append(json.loads(r["payload"]).get("text", ""))
    out = sorted(
        [{"id": c["id"], "label": c["label"], "count": len(by_cluster[c["id"]]),
          "items": by_cluster[c["id"]]} for c in clusters],
        key=lambda x: -x["count"],
    )
    return {"type": "opentext", "n": n, "clustered": True, "clusters": out}


def _materialize_text_clusters(conn, run_id: str, slide_id: str, result: dict) -> None:
    """Materializza il clustering a un asse (open text): cluster kind='theme' + response.cluster_id."""
    now = db.now_iso()
    conn.execute("DELETE FROM cluster WHERE run_id=? AND slide_id=?", (run_id, slide_id))
    conn.execute(
        "UPDATE response SET cluster_id=NULL WHERE run_id=? AND slide_id=?", (run_id, slide_id)
    )
    cmap: dict = {}
    for i, c in enumerate(result.get("clusters", [])):
        cid = db.new_id()
        conn.execute(
            "INSERT INTO cluster (id, run_id, slide_id, kind, label, ord, generated_at) "
            "VALUES (?,?,?,'theme',?,?,?)",
            (cid, run_id, slide_id, c.get("label", "—"), i, now),
        )
        cmap[c.get("id")] = cid
    resp_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM response WHERE run_id=? AND slide_id=? AND status='visible' "
            "ORDER BY created_at",
            (run_id, slide_id),
        ).fetchall()
    ]
    for a in result.get("assignments", []):
        nn = a.get("n")
        if not isinstance(nn, int) or nn < 1 or nn > len(resp_ids):
            continue
        conn.execute(
            "UPDATE response SET cluster_id=? WHERE id=?",
            (cmap.get(a.get("cluster")), resp_ids[nn - 1]),
        )


def _materialize_clusters(conn, run_id: str, slide_id: str, result: dict) -> None:
    """Salva l'esito del clustering LLM: cancella i precedenti, crea cluster, assegna le risposte."""
    now = db.now_iso()
    conn.execute("DELETE FROM cluster WHERE run_id=? AND slide_id=?", (run_id, slide_id))
    conn.execute(
        "UPDATE response SET claim_cluster_id=NULL, arg_cluster_id=NULL "
        "WHERE run_id=? AND slide_id=?",
        (run_id, slide_id),
    )
    claim_map: dict = {}
    for i, c in enumerate(result.get("claim_clusters", [])):
        cid = db.new_id()
        conn.execute(
            "INSERT INTO cluster (id, run_id, slide_id, kind, label, ord, generated_at) "
            "VALUES (?,?,?,'claim',?,?,?)",
            (cid, run_id, slide_id, c.get("label", "—"), i, now),
        )
        claim_map[c.get("id")] = cid
    arg_map: dict = {}
    for i, a in enumerate(result.get("arg_clusters", [])):
        aid = db.new_id()
        conn.execute(
            "INSERT INTO cluster (id, run_id, slide_id, kind, label, ord, generated_at) "
            "VALUES (?,?,?,'arg',?,?,?)",
            (aid, run_id, slide_id, a.get("label", "—"), i, now),
        )
        arg_map[a.get("id")] = aid
    resp_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM response WHERE run_id=? AND slide_id=? AND status='visible' "
            "ORDER BY created_at",
            (run_id, slide_id),
        ).fetchall()
    ]
    for a in result.get("assignments", []):
        n = a.get("n")
        if not isinstance(n, int) or n < 1 or n > len(resp_ids):
            continue
        conn.execute(
            "UPDATE response SET claim_cluster_id=?, arg_cluster_id=? WHERE id=?",
            (claim_map.get(a.get("claim")), arg_map.get(a.get("arg")), resp_ids[n - 1]),
        )


def _qa_results(conn, run_id: str, slide_id: str) -> dict:
    """qa ha bisogno del DB per i conteggi voti → non passa per aggregate()."""
    rows = conn.execute(
        "SELECT r.id, r.payload, "
        "  (SELECT COUNT(*) FROM qa_vote v WHERE v.response_id=r.id) AS votes "
        "FROM response r WHERE r.run_id=? AND r.slide_id=? AND r.status='visible' "
        "ORDER BY votes DESC, r.created_at ASC",
        (run_id, slide_id),
    ).fetchall()
    items = [
        {"id": r["id"], "text": json.loads(r["payload"]).get("text", ""), "votes": r["votes"]}
        for r in rows
    ]
    return {"type": "qa", "n": len(items), "items": items}


def _moderation(conn, run_id: str, slide_id: str) -> list:
    rows = conn.execute(
        "SELECT id, payload, status, created_at FROM response "
        "WHERE run_id=? AND slide_id=? ORDER BY created_at DESC",
        (run_id, slide_id),
    ).fetchall()
    out = []
    for r in rows:
        p = json.loads(r["payload"])
        # argstep: si modera cio che questa tappa ha raccolto (l'obiezione, non il claim ereditato)
        text = p.get("objection") or p.get("text") or p.get("claim") or ""
        out.append(
            {
                "id": r["id"],
                "text": text,
                "justification": p.get("claim") if p.get("objection") else p.get("justification"),
                "status": r["status"],
            }
        )
    return out


# ----------------------------------------------------------------------------
# API — presenter
# ----------------------------------------------------------------------------
class PresentationIn(BaseModel):
    title: str
    owner: str = "spit"


def _check_owner(conn, pid: str, user: dict):
    """Alza 404/403 se la presentation non esiste o non è dell'utente."""
    row = conn.execute("SELECT owner FROM presentation WHERE id=?", (pid,)).fetchone()
    if not row:
        raise HTTPException(404, "presentation not found")
    if row["owner"] != user["id"]:
        raise HTTPException(403, "non autorizzato")


def _pid_of_run(conn, rid: str) -> str:
    r = conn.execute("SELECT presentation_id FROM run WHERE id=?", (rid,)).fetchone()
    if not r:
        raise HTTPException(404, "run not found")
    return r["presentation_id"]


def _pid_of_slide(conn, sid: str) -> str:
    r = conn.execute("SELECT presentation_id FROM slide WHERE id=?", (sid,)).fetchone()
    if not r:
        raise HTTPException(404, "slide not found")
    return r["presentation_id"]


@app.get("/api/presentations")
def list_presentations(user: dict = CurrentUser):
    with db.get_conn() as conn:
        rows = conn.execute(
            "SELECT p.id, p.title, p.join_code, "
            "  (SELECT COUNT(*) FROM slide s WHERE s.presentation_id=p.id) AS n_slides "
            "FROM presentation p WHERE p.owner=? ORDER BY p.created_at DESC",
            (user["id"],),
        ).fetchall()
        return [dict(r) for r in rows]


class RenameIn(BaseModel):
    title: str


@app.patch("/api/presentations/{pid}")
def rename_presentation(pid: str, body: RenameIn, user: dict = CurrentUser):
    title = body.title.strip()
    if not title:
        raise HTTPException(400, "titolo vuoto")
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        conn.execute("UPDATE presentation SET title=? WHERE id=?", (title, pid))
        return {"ok": True, "title": title}


@app.post("/api/presentations")
def create_presentation(body: PresentationIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        pid = db.new_id()
        code = db.new_join_code(conn)
        conn.execute(
            "INSERT INTO presentation (id, title, owner, join_code, created_at) "
            "VALUES (?,?,?,?,?)",
            (pid, body.title, user["id"], code, db.now_iso()),
        )
        return {"id": pid, "title": body.title, "join_code": code}


def _validate_argstep(conn, pid: str, cfg: dict, self_id: str | None = None) -> None:
    """La tappa N puo ereditare solo da una slide argstep di questa deck con esattamente
    N-1 campi. Senza questo vincolo il carryover produrrebbe campi vuoti a runtime."""
    n = len(_argstep_fields(cfg))
    src = cfg.get("carry_from")
    if n == 1:
        if src:
            raise HTTPException(400, "la prima tappa non eredita nulla")
        return
    if not src:
        raise HTTPException(400, "tappa oltre la prima: serve la slide da cui ereditare")
    if src == self_id:
        raise HTTPException(400, "una tappa non puo ereditare da se stessa")
    row = conn.execute(
        "SELECT type, config FROM slide WHERE id=? AND presentation_id=?", (src, pid)
    ).fetchone()
    if not row or row["type"] != "argstep":
        raise HTTPException(400, "la slide sorgente non esiste o non e una tappa argomentativa")
    if len(_argstep_fields(json.loads(row["config"]))) != n - 1:
        raise HTTPException(400, f"la sorgente deve avere {n - 1} campi, non di piu ne di meno")


class SlideIn(BaseModel):
    type: str
    question: str
    config: dict = {}
    pair_id: str | None = None   # pre/post: questa slide è la POST della slide indicata (PRE)
    presenter_notes: str = ""    # visibili solo al presenter


@app.post("/api/presentations/{pid}/slides")
def add_slide(pid: str, body: SlideIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        if body.pair_id:
            tgt = conn.execute(
                "SELECT type FROM slide WHERE id=? AND presentation_id=?",
                (body.pair_id, pid),
            ).fetchone()
            if not tgt:
                raise HTTPException(400, "slide pre/post non valida")
            if body.type not in ("scale", "mc", "quadrant") or tgt["type"] != body.type:
                raise HTTPException(
                    400, "pre/post consentito solo tra slide dello stesso tipo (scale/mc/quadrant)"
                )
        if body.type == "argstep":
            _validate_argstep(conn, pid, body.config)
        ord_row = conn.execute(
            "SELECT COALESCE(MAX(ord), 0) + 1 AS n FROM slide WHERE presentation_id=?",
            (pid,),
        ).fetchone()
        sid = db.new_id()
        conn.execute(
            "INSERT INTO slide (id, presentation_id, ord, type, question, config, pair_id, presenter_notes) "
            "VALUES (?,?,?,?,?,?,?,?)",
            (sid, pid, ord_row["n"], body.type, body.question,
             json.dumps(body.config), body.pair_id, body.presenter_notes),
        )
        return {"id": sid, "ord": ord_row["n"]}


class SlideEdit(BaseModel):
    edit_structure: bool = False       # False → aggiorna solo le note (tier "sempre"); True → struttura (gated)
    question: str | None = None
    config: dict | None = None
    pair_id: str | None = None         # con edit_structure=True: None = scollega la coppia
    presenter_notes: str | None = None


@app.patch("/api/slides/{slide_id}")
def edit_slide(slide_id: str, body: SlideEdit, user: dict = CurrentUser):
    """Modifica una slide. Le note sono sempre editabili; domanda/config/pair solo se la slide
    non ha ancora risposte (gate per-slide, lato server)."""
    with db.get_conn() as conn:
        pid = _pid_of_slide(conn, slide_id)
        _check_owner(conn, pid, user)
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (slide_id,)).fetchone()

        if not body.edit_structure:
            # tier "sempre": solo presenter notes, nessun gate
            if body.presenter_notes is not None:
                conn.execute(
                    "UPDATE slide SET presenter_notes=? WHERE id=?",
                    (body.presenter_notes, slide_id),
                )
            return {"ok": True, "scope": "notes"}

        # tier strutturale: vietato se la slide ha risposte
        if conn.execute(
            "SELECT 1 FROM response WHERE slide_id=? LIMIT 1", (slide_id,)
        ).fetchone():
            raise HTTPException(409, "la slide ha già risposte: modifica strutturale non permessa")

        pair = body.pair_id or None
        if pair:
            if pair == slide_id:
                raise HTTPException(400, "una slide non può essere POST di se stessa")
            tgt = conn.execute(
                "SELECT type FROM slide WHERE id=? AND presentation_id=?", (pair, pid)
            ).fetchone()
            if not tgt or slide["type"] not in ("scale", "mc", "quadrant") or tgt["type"] != slide["type"]:
                raise HTTPException(
                    400, "pre/post consentito solo tra slide dello stesso tipo (scale/mc/quadrant)"
                )
        if slide["type"] == "argstep":
            _validate_argstep(
                conn, pid,
                body.config if body.config is not None else json.loads(slide["config"]),
                slide_id,
            )
        q = body.question if body.question is not None else slide["question"]
        cfg = json.dumps(body.config) if body.config is not None else slide["config"]
        notes = body.presenter_notes if body.presenter_notes is not None else slide["presenter_notes"]
        conn.execute(
            "UPDATE slide SET question=?, config=?, pair_id=?, presenter_notes=? WHERE id=?",
            (q, cfg, pair, notes, slide_id),
        )
        return {"ok": True, "scope": "full"}


class RunIn(BaseModel):
    label: str | None = None


@app.post("/api/presentations/{pid}/runs")
def start_run(pid: str, body: RunIn, user: dict = CurrentUser):
    """Crea un nuovo run e lo imposta come run corrente della presentation."""
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        rid = db.new_id()
        conn.execute(
            "INSERT INTO run (id, presentation_id, label, started_at) VALUES (?,?,?,?)",
            (rid, pid, body.label, db.now_iso()),
        )
        conn.execute(
            "UPDATE presentation SET active_run_id=? WHERE id=?", (rid, pid)
        )
        # la presenza dei run precedenti non serve piu' a nessuno: e' un heartbeat, non storia
        conn.execute(
            "DELETE FROM presence WHERE run_id IN "
            "(SELECT id FROM run WHERE presentation_id=? AND id<>?)",
            (pid, rid),
        )
        return {"run_id": rid}


@app.delete("/api/presentations/{pid}/runs")
def purge_runs(pid: str, user: dict = CurrentUser):
    """Distrugge TUTTI i run della deck (risposte, voti, moderazione, cluster) — tiene slide e deck."""
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        run_ids = [
            r["id"] for r in conn.execute(
                "SELECT id FROM run WHERE presentation_id=?", (pid,)
            ).fetchall()
        ]
        for rid in run_ids:
            conn.execute(
                "DELETE FROM qa_vote WHERE response_id IN (SELECT id FROM response WHERE run_id=?)",
                (rid,),
            )
            conn.execute("DELETE FROM response WHERE run_id=?", (rid,))
            conn.execute("DELETE FROM run_slide WHERE run_id=?", (rid,))
            conn.execute("DELETE FROM mc_option WHERE run_id=?", (rid,))
            conn.execute("DELETE FROM cluster WHERE run_id=?", (rid,))
            conn.execute("DELETE FROM carry_assign WHERE run_id=?", (rid,))
            conn.execute("DELETE FROM presence WHERE run_id=?", (rid,))
        conn.execute("UPDATE presentation SET active_run_id=NULL WHERE id=?", (pid,))
        conn.execute("DELETE FROM run WHERE presentation_id=?", (pid,))
        return {"ok": True, "purged_runs": len(run_ids)}


class ActivateIn(BaseModel):
    slide_id: str


@app.post("/api/runs/{rid}/activate")
def activate_slide(rid: str, body: ActivateIn, user: dict = CurrentUser):
    """Rende attiva una slide nel run e la apre al voto."""
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        run = conn.execute("SELECT * FROM run WHERE id=?", (rid,)).fetchone()
        if not run:
            raise HTTPException(404, "run not found")
        slide = conn.execute(
            "SELECT * FROM slide WHERE id=? AND presentation_id=?",
            (body.slide_id, run["presentation_id"]),
        ).fetchone()
        if not slide:
            raise HTTPException(404, "slide not in this presentation")
        conn.execute("UPDATE run SET active_slide_id=? WHERE id=?", (body.slide_id, rid))
        conn.execute(
            "INSERT INTO run_slide (run_id, slide_id, state) VALUES (?,?, 'open') "
            "ON CONFLICT(run_id, slide_id) DO UPDATE SET state='open'",
            (rid, body.slide_id),
        )
        if slide["type"] == "timer" and _timer_config(slide)["autostart"]:
            _timer_start(conn, rid, slide, only_if_absent=True)
        return {"active_slide_id": body.slide_id, "state": "open"}


class StateIn(BaseModel):
    slide_id: str
    state: str  # open | closed | revealed


@app.post("/api/runs/{rid}/state")
def set_state(rid: str, body: StateIn, user: dict = CurrentUser):
    if body.state not in ("open", "closed", "revealed"):
        raise HTTPException(400, "invalid state")
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        conn.execute(
            "INSERT INTO run_slide (run_id, slide_id, state) VALUES (?,?,?) "
            "ON CONFLICT(run_id, slide_id) DO UPDATE SET state=excluded.state",
            (rid, body.slide_id, body.state),
        )
        return {"slide_id": body.slide_id, "state": body.state}


@app.get("/api/presentations/{pid}")
def presenter_view(pid: str, user: dict = CurrentUser):
    """Vista completa per il presenter: deck + slide + run corrente + risultati attivi."""
    with db.get_conn() as conn:
        pres = conn.execute("SELECT * FROM presentation WHERE id=?", (pid,)).fetchone()
        if not pres:
            raise HTTPException(404, "presentation not found")
        if pres["owner"] != user["id"]:
            raise HTTPException(403, "non autorizzato")
        slides = conn.execute(
            "SELECT * FROM slide WHERE presentation_id=? ORDER BY ord", (pid,)
        ).fetchall()
        responded = {
            r["slide_id"] for r in conn.execute(
                "SELECT DISTINCT slide_id FROM response "
                "WHERE run_id IN (SELECT id FROM run WHERE presentation_id=?)", (pid,)
            ).fetchall()
        }

        def _sd(s):
            d = _slide_dict(s)
            d["has_responses"] = s["id"] in responded  # gate editing strutturale (per-slide)
            return d

        out = {
            "id": pres["id"],
            "title": pres["title"],
            "join_code": pres["join_code"],
            "active_run_id": pres["active_run_id"],
            "slides": [_sd(s) for s in slides],
            "run": None,
        }
        rid = pres["active_run_id"]
        if rid:
            run = conn.execute("SELECT * FROM run WHERE id=?", (rid,)).fetchone()
            states = {
                r["slide_id"]: r["state"]
                for r in conn.execute(
                    "SELECT slide_id, state FROM run_slide WHERE run_id=?", (rid,)
                ).fetchall()
            }
            active = run["active_slide_id"]
            results = None
            moderation = None
            pair = None
            aslide = conn.execute("SELECT * FROM slide WHERE id=?", (active,)).fetchone() if active else None
            if aslide is None:
                active = None  # puntatore penzolante (slide cancellata) → nessuna attiva
            if active:
                results = _results(conn, rid, aslide)
                if aslide["type"] in MODERATED_TYPES:
                    moderation = _moderation(conn, rid, active)
                if aslide["pair_id"]:
                    pslide = conn.execute(
                        "SELECT * FROM slide WHERE id=?", (aslide["pair_id"],)
                    ).fetchone()
                    if pslide:
                        pair = {
                            "slide": _slide_dict(pslide),
                            "results": _results(conn, rid, pslide),
                        }
            out["run"] = {
                "id": rid,
                "label": run["label"],
                "active_slide_id": active,
                "states": states,
                "results": results,
                "moderation": moderation,
                "pair": pair,
                "present": _present_count(conn, rid),
                "voters": _voter_count(conn, rid, active) if active else 0,
            }
            if active and aslide and aslide["type"] == "moonshot":
                _ensure_moonshot_lobby(conn, rid, active)
                out["run"]["moonshot"] = _moonshot_state(conn, rid, aslide)
            if active and aslide and aslide["type"] == "timer":
                out["run"]["timer"] = _timer_state(conn, rid, aslide)
        return out


class SlideStatusIn(BaseModel):
    status: str  # visible | hidden | flagged


@app.post("/api/responses/{response_id}/status")
def set_response_status(response_id: str, body: SlideStatusIn, user: dict = CurrentUser):
    if body.status not in ("visible", "hidden", "flagged"):
        raise HTTPException(400, "invalid status")
    with db.get_conn() as conn:
        r = conn.execute("SELECT run_id FROM response WHERE id=?", (response_id,)).fetchone()
        if not r:
            raise HTTPException(404, "response not found")
        _check_owner(conn, _pid_of_run(conn, r["run_id"]), user)
        conn.execute("UPDATE response SET status=? WHERE id=?", (body.status, response_id))
        return {"id": response_id, "status": body.status}


@app.delete("/api/slides/{slide_id}")
def delete_slide(slide_id: str, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_slide(conn, slide_id), user)
        # azzera i run che avevano questa slide come attiva (evita puntatori penzolanti)
        conn.execute("UPDATE run SET active_slide_id=NULL WHERE active_slide_id=?", (slide_id,))
        conn.execute(
            "DELETE FROM qa_vote WHERE response_id IN (SELECT id FROM response WHERE slide_id=?)",
            (slide_id,),
        )
        conn.execute("DELETE FROM mc_option WHERE slide_id=?", (slide_id,))
        conn.execute("DELETE FROM cluster WHERE slide_id=?", (slide_id,))
        conn.execute("DELETE FROM response WHERE slide_id=?", (slide_id,))
        conn.execute("DELETE FROM run_slide WHERE slide_id=?", (slide_id,))
        conn.execute("DELETE FROM carry_assign WHERE slide_id=?", (slide_id,))
        # scollega eventuali coppie pre/post che puntavano a questa slide
        conn.execute("UPDATE slide SET pair_id=NULL WHERE pair_id=?", (slide_id,))
        # e le tappe argstep che ereditavano da questa: meglio scollegate che penzolanti
        for row in conn.execute(
            "SELECT id, config FROM slide WHERE presentation_id=? AND type='argstep'",
            (_pid_of_slide(conn, slide_id),),
        ).fetchall():
            cfg = json.loads(row["config"])
            if cfg.get("carry_from") == slide_id:
                cfg["carry_from"] = None
                conn.execute("UPDATE slide SET config=? WHERE id=?", (json.dumps(cfg), row["id"]))
        conn.execute("DELETE FROM slide WHERE id=?", (slide_id,))
        return {"ok": True}


def _purge_presentation(conn, pid: str) -> tuple[int, int]:
    """Distrugge una deck e tutto cio che vi pende: slide, run, risposte, voti, cluster,
    assegnazioni peer. Condivisa fra la cancellazione della deck e quella dell'utente,
    cosi le due non possono divergere e lasciare orfani."""
    slide_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM slide WHERE presentation_id=?", (pid,)
        ).fetchall()
    ]
    run_ids = [
        r["id"] for r in conn.execute(
            "SELECT id FROM run WHERE presentation_id=?", (pid,)
        ).fetchall()
    ]
    for rid in run_ids:
        conn.execute(
            "DELETE FROM qa_vote WHERE response_id IN (SELECT id FROM response WHERE run_id=?)",
            (rid,),
        )
        conn.execute("DELETE FROM response WHERE run_id=?", (rid,))
        conn.execute("DELETE FROM run_slide WHERE run_id=?", (rid,))
        conn.execute("DELETE FROM mc_option WHERE run_id=?", (rid,))
        conn.execute("DELETE FROM cluster WHERE run_id=?", (rid,))
        conn.execute("DELETE FROM carry_assign WHERE run_id=?", (rid,))
        conn.execute("DELETE FROM presence WHERE run_id=?", (rid,))
    conn.execute("DELETE FROM run WHERE presentation_id=?", (pid,))
    for sid in slide_ids:
        conn.execute("DELETE FROM slide WHERE id=?", (sid,))
    conn.execute("DELETE FROM presentation WHERE id=?", (pid,))
    return len(slide_ids), len(run_ids)


@app.delete("/api/presentations/{pid}")
def delete_presentation(pid: str, user: dict = CurrentUser):
    """Cancella una deck e TUTTI i dati associati (slide, run, risposte, voti)."""
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        n_slides, n_runs = _purge_presentation(conn, pid)
        return {"ok": True, "deleted_slides": n_slides, "deleted_runs": n_runs}


class ReorderIn(BaseModel):
    slide_ids: list[str]


@app.post("/api/presentations/{pid}/reorder")
def reorder_slides(pid: str, body: ReorderIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        existing = {
            r["id"] for r in conn.execute(
                "SELECT id FROM slide WHERE presentation_id=?", (pid,)
            ).fetchall()
        }
        if set(body.slide_ids) != existing:
            raise HTTPException(400, "lista slide incompleta")
        # due fasi per non violare UNIQUE(presentation_id, ord)
        for sid in body.slide_ids:
            conn.execute("UPDATE slide SET ord=ord+10000 WHERE id=?", (sid,))
        for i, sid in enumerate(body.slide_ids, start=1):
            conn.execute("UPDATE slide SET ord=? WHERE id=?", (i, sid))
        return {"ok": True}


# ── Export / Import deck (solo config, non i dati delle run) ──────────────────
@app.get("/api/presentations/{pid}/export")
def export_deck(pid: str, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        p = conn.execute("SELECT * FROM presentation WHERE id=?", (pid,)).fetchone()
        slides = conn.execute(
            "SELECT * FROM slide WHERE presentation_id=? ORDER BY ord", (pid,)
        ).fetchall()
        return {
            "roompulse_deck": 1,
            "title": p["title"],
            "slides": [
                {
                    "ref": s["id"],
                    "type": s["type"],
                    "question": s["question"],
                    "config": json.loads(s["config"]),
                    "pair_ref": s["pair_id"],
                    "presenter_notes": s["presenter_notes"],
                }
                for s in slides
            ],
        }


class ImportSlide(BaseModel):
    ref: str | None = None
    type: str
    question: str
    config: dict = {}
    pair_ref: str | None = None
    presenter_notes: str = ""


class ImportDeck(BaseModel):
    title: str
    slides: list[ImportSlide]


@app.post("/api/presentations/import")
def import_deck(body: ImportDeck, user: dict = CurrentUser):
    with db.get_conn() as conn:
        pid = db.new_id()
        code = db.new_join_code(conn)
        conn.execute(
            "INSERT INTO presentation (id, title, owner, join_code, created_at) VALUES (?,?,?,?,?)",
            (pid, body.title, user["id"], code, db.now_iso()),
        )
        new_ids: list[str] = []
        refmap: dict[str, str] = {}
        for i, s in enumerate(body.slides, start=1):
            sid = db.new_id()
            conn.execute(
                "INSERT INTO slide (id, presentation_id, ord, type, question, config, presenter_notes) "
                "VALUES (?,?,?,?,?,?,?)",
                (sid, pid, i, s.type, s.question, json.dumps(s.config), s.presenter_notes),
            )
            new_ids.append(sid)
            if s.ref:
                refmap[s.ref] = sid
        for i, s in enumerate(body.slides):
            if s.pair_ref and s.pair_ref in refmap:
                conn.execute(
                    "UPDATE slide SET pair_id=? WHERE id=?", (refmap[s.pair_ref], new_ids[i])
                )
            if s.type == "argstep":
                _remap_carry(conn, new_ids[i], json.dumps(s.config), refmap)
        return {"id": pid, "title": body.title, "join_code": code, "n_slides": len(body.slides)}


def _fmt_answer(stype: str, p: dict, optmap: dict) -> str:
    """Formatta il payload di una risposta in una stringa leggibile per il CSV."""
    if stype == "mc":
        picks = p["choices"] if isinstance(p.get("choices"), list) else (
            [p["choice"]] if p.get("choice") is not None else [])
        return "; ".join(optmap.get(c, c) for c in picks)
    if stype == "scale":
        return str(p.get("value", ""))
    if stype == "quadrant":
        return f'{p.get("x", "")},{p.get("y", "")}'
    if stype == "ranking":
        return " > ".join(optmap.get(i, i) for i in p.get("order", []))
    if stype == "points":
        return "; ".join(f"{optmap.get(k, k)}:{v}" for k, v in (p.get("alloc") or {}).items())
    if stype in ("wordcloud", "opentext", "qa"):
        return p.get("text", "")
    if stype == "groups":
        return p.get("group_name", "")
    if stype == "donut":
        return str(p.get("score", ""))
    if stype in ("argpoll", "argstep"):
        return ""  # claim/justification/objection vanno nelle loro colonne
    return json.dumps(p, ensure_ascii=False)


_EXPORT_HEADER = ["slide_ord", "slide_type", "question", "participant", "created_at",
                  "status", "claim", "justification", "objection", "answer", "cluster",
                  "arg_cluster", "carry_source"]


def _run_export_rows(conn, rid: str, slides) -> list[list]:
    """Righe grezze di un run (una per risposta), con label risolte e cluster."""
    clabels = {
        c["id"]: c["label"]
        for c in conn.execute("SELECT id, label FROM cluster WHERE run_id=?", (rid,)).fetchall()
    }
    out: list[list] = []
    for s in slides:
        cfg = json.loads(s["config"])
        optmap: dict = {}
        if s["type"] == "mc":
            for o in cfg.get("options", []):
                optmap[o["id"]] = o["label"]
            for mo in conn.execute(
                "SELECT id, label FROM mc_option WHERE run_id=? AND slide_id=?", (rid, s["id"])
            ).fetchall():
                optmap[mo["id"]] = mo["label"]
        elif s["type"] == "ranking":
            for o in cfg.get("items", []):
                optmap[o["id"]] = o["label"]
        elif s["type"] == "points":
            for o in cfg.get("options", []):
                optmap[o["id"]] = o["label"]
        rows = conn.execute(
            "SELECT participant_token, payload, status, created_at, "
            "claim_cluster_id, arg_cluster_id, cluster_id, source_response_id FROM response "
            "WHERE run_id=? AND slide_id=? ORDER BY created_at",
            (rid, s["id"]),
        ).fetchall()
        for r in rows:
            p = json.loads(r["payload"])
            argy = s["type"] in ("argpoll", "argstep")
            claim = p.get("claim", "") if argy else ""
            just = p.get("justification", "") if argy else ""
            obj = p.get("objection", "") if argy else ""
            primary = clabels.get(r["claim_cluster_id"]) or clabels.get(r["cluster_id"]) or ""
            argc = clabels.get(r["arg_cluster_id"]) or ""
            out.append([s["ord"], s["type"], s["question"], r["participant_token"],
                        r["created_at"], r["status"], claim, just, obj,
                        _fmt_answer(s["type"], p, optmap), primary, argc,
                        r["source_response_id"] or ""])
    return out


def _safe_sheet_name(base: str, used: set) -> str:
    """Nome foglio Excel valido: caratteri vietati rimossi, ≤31 char, unico."""
    name = re.sub(r"[:\\/?*\[\]]", " ", base or "").strip()[:31] or "Run"
    stem, i = name, 2
    while name.lower() in used:
        suffix = f" ({i})"
        name = stem[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


@app.get("/api/presentations/{pid}/data.xlsx")
def export_data(pid: str, user: dict = CurrentUser):
    """Workbook Excel: un foglio Overview (elenco run) + un foglio per run con le risposte grezze."""
    with db.get_conn() as conn:
        _check_owner(conn, pid, user)
        pres = conn.execute("SELECT * FROM presentation WHERE id=?", (pid,)).fetchone()
        slides = conn.execute(
            "SELECT * FROM slide WHERE presentation_id=? ORDER BY ord", (pid,)
        ).fetchall()
        runs = conn.execute(
            "SELECT * FROM run WHERE presentation_id=? ORDER BY started_at", (pid,)
        ).fetchall()
        if not runs:
            raise HTTPException(400, "nessun run da esportare")

        wb = Workbook()
        ov = wb.active
        ov.title = "Overview"
        ov.append(["run", "label", "started_at", "responses", "slides answered"])
        used = {"overview"}

        for i, run in enumerate(runs, start=1):
            rows = _run_export_rows(conn, run["id"], slides)
            answered = len({r[0] for r in rows})  # slide_ord distinti con risposte
            ov.append([i, run["label"] or "", run["started_at"], len(rows), answered])
            if not rows:
                continue  # run vuoto: resta in Overview, ma niente foglio dedicato
            base = run["label"] or (run["started_at"] or "")[:16].replace("T", " ")
            ws = wb.create_sheet(_safe_sheet_name(base, used))
            ws.append(_EXPORT_HEADER)
            for r in rows:
                ws.append(r)

        buf = io.BytesIO()
        wb.save(buf)
        fname = "".join(ch if ch.isalnum() else "_" for ch in (pres["title"] or "deck")).lower()[:40]
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname}_dati.xlsx"'},
        )


# ----------------------------------------------------------------------------
# API — audience (pubblico, via join_code)
# ----------------------------------------------------------------------------
def _resolve_run(conn, code: str):
    pres = conn.execute(
        "SELECT * FROM presentation WHERE join_code=?", (code,)
    ).fetchone()
    if not pres:
        raise HTTPException(404, "codice non valido")
    return pres


@app.get("/api/live/{code}")
def live(code: str, t: str | None = None):
    """Ciò che il client del pubblico richiede in polling. `t` = token (personalizza moonshot)."""
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            return {"status": "waiting", "title": pres["title"]}
        run = conn.execute("SELECT * FROM run WHERE id=?", (rid,)).fetchone()
        active = run["active_slide_id"]
        if not active:
            return {"status": "waiting", "title": pres["title"]}
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (active,)).fetchone()
        if slide is None:  # puntatore penzolante
            return {"status": "waiting", "title": pres["title"]}
        if t:
            _touch_presence(conn, rid, t)
        state = _run_slide_state(conn, rid, active)
        payload = {
            "status": "live",
            "run_id": rid,
            "state": state,
            "slide": _slide_dict(slide),
        }
        payload["slide"].pop("presenter_notes", None)  # mai al pubblico
        if slide["type"] == "qa":
            # qa: il pubblico vede e vota le domande altrui anche mentre è 'open'
            payload["feed"] = _qa_results(conn, rid, active)
        if slide["type"] == "donut":
            # leaderboard non segreta: sempre allegata (l'audience mostra top3 se open, tutto se closed)
            payload["results"] = _results(conn, rid, slide)
        if slide["type"] == "moonshot":
            _ensure_moonshot_lobby(conn, rid, active)
            payload["moonshot"] = _moonshot_state(conn, rid, slide, t)
        if slide["type"] == "timer":
            payload["timer"] = _timer_state(conn, rid, slide)
        if slide["type"] == "argstep":
            cfg = payload["slide"]["config"]
            cfg["fields"] = _argstep_fields(cfg)
            carry = _resolve_carry(conn, rid, slide, t)
            if carry:
                payload["carry"] = carry
            elif cfg.get("carry_from"):
                # arrivato tardi (o nessun peer disponibile): i campi ereditati
                # diventano scrivibili invece di bloccare la partecipazione
                payload["carry_missing"] = True
        if slide["type"] == "mc":
            # mc: opzioni base + quelle aggiunte dai partecipanti (allow_other)
            cfg = payload["slide"]["config"]
            cfg["options"] = _merged_mc_options(conn, rid, slide)
            cfg.pop("correct", None)  # mai svelare la risposta corretta durante il voto
        if state == "revealed":
            payload["results"] = _results(conn, rid, slide)
            if slide["pair_id"]:  # pre/post: l'audience vede lo stesso confronto del presenter
                pslide = conn.execute(
                    "SELECT * FROM slide WHERE id=?", (slide["pair_id"],)
                ).fetchone()
                if pslide:
                    ps = _slide_dict(pslide)
                    ps.pop("presenter_notes", None)
                    payload["pair"] = {
                        "slide": ps,
                        "results": _results(conn, rid, pslide),
                    }
        return payload


class AddOptionIn(BaseModel):
    slide_id: str
    label: str


@app.post("/api/live/{code}/add-option")
def add_option(code: str, body: AddOptionIn):
    """Un partecipante aggiunge un'opzione mc (se allow_other) — visibile a tutti."""
    label = body.label.strip()
    if not label:
        raise HTTPException(400, "etichetta vuota")
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] != "mc":
            raise HTTPException(404, "slide non valida")
        if not json.loads(slide["config"]).get("allow_other"):
            raise HTTPException(403, "aggiunta opzioni non consentita")
        oid = "x" + db.new_id()[:8]
        conn.execute(
            "INSERT INTO mc_option (id, run_id, slide_id, label, created_at) VALUES (?,?,?,?,?)",
            (oid, rid, body.slide_id, label, db.now_iso()),
        )
        return {"id": oid, "label": label}


class UpvoteIn(BaseModel):
    response_id: str
    token: str


@app.post("/api/live/{code}/upvote")
def upvote(code: str, body: UpvoteIn):
    """Toggle dell'upvote di una domanda qa (un voto per token)."""
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        row = conn.execute(
            "SELECT slide_id FROM response WHERE id=?", (body.response_id,)
        ).fetchone()
        if not row:
            raise HTTPException(404, "domanda non trovata")
        if _run_slide_state(conn, rid, row["slide_id"]) != "open":
            raise HTTPException(409, "domande chiuse")
        exists = conn.execute(
            "SELECT 1 FROM qa_vote WHERE response_id=? AND token=?",
            (body.response_id, body.token),
        ).fetchone()
        if exists:
            conn.execute(
                "DELETE FROM qa_vote WHERE response_id=? AND token=?",
                (body.response_id, body.token),
            )
            return {"voted": False}
        conn.execute(
            "INSERT INTO qa_vote (run_id, response_id, token) VALUES (?,?,?)",
            (rid, body.response_id, body.token),
        )
        return {"voted": True}


class RespondIn(BaseModel):
    slide_id: str
    token: str
    payload: dict


@app.post("/api/live/{code}/respond")
def respond(code: str, body: RespondIn):
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        run = conn.execute("SELECT * FROM run WHERE id=?", (rid,)).fetchone()
        if run["active_slide_id"] != body.slide_id:
            raise HTTPException(409, "la slide non è più attiva")
        if _run_slide_state(conn, rid, body.slide_id) != "open":
            raise HTTPException(409, "voto chiuso")
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()

        # quiz mc: risposta definitiva (non sovrascrivibile) + ritorno la correttezza
        if slide["type"] == "mc":
            cfg = json.loads(slide["config"])
            if cfg.get("quiz"):
                correct = set(cfg.get("correct", []))

                def _picks(p):
                    return (p.get("choices") if isinstance(p.get("choices"), list)
                            else ([p["choice"]] if p.get("choice") is not None else []))

                ex = conn.execute(
                    "SELECT payload FROM response WHERE run_id=? AND slide_id=? AND participant_token=?",
                    (rid, body.slide_id, body.token),
                ).fetchone()
                if ex:  # ha già risposto → bloccato, ritorno il suo esito originale
                    picks = _picks(json.loads(ex["payload"]))
                    return {"ok": True, "locked": True,
                            "quiz": {"correct": set(picks) == correct, "correct_ids": list(correct)}}
                conn.execute(
                    "INSERT INTO response (id, run_id, slide_id, participant_token, payload, created_at) "
                    "VALUES (?,?,?,?,?,?)",
                    (db.new_id(), rid, body.slide_id, body.token, json.dumps(body.payload), db.now_iso()),
                )
                picks = _picks(body.payload)
                return {"ok": True, "quiz": {"correct": set(picks) == correct, "correct_ids": list(correct)}}

        # donut: best-score per token (tieni il massimo) + cap anti-assurdo
        if slide["type"] == "donut":
            raw = body.payload.get("score")
            if not isinstance(raw, (int, float)) or raw < 0 or raw > 100000:
                raise HTTPException(400, "punteggio non valido")
            score = int(raw)
            name = str(body.payload.get("name") or "").strip()[:40]
            ex = conn.execute(
                "SELECT id, payload FROM response WHERE run_id=? AND slide_id=? AND participant_token=?",
                (rid, body.slide_id, body.token),
            ).fetchone()
            if ex:
                prev = json.loads(ex["payload"])
                payload = {"score": max(score, int(prev.get("score", 0))), "name": name or prev.get("name", "—")}
                conn.execute("UPDATE response SET payload=? WHERE id=?", (json.dumps(payload), ex["id"]))
            else:
                payload = {"score": score, "name": name or "—"}
                conn.execute(
                    "INSERT INTO response (id, run_id, slide_id, participant_token, payload, created_at) "
                    "VALUES (?,?,?,?,?,?)",
                    (db.new_id(), rid, body.slide_id, body.token, json.dumps(payload), db.now_iso()),
                )
            return {"ok": True, "best": payload["score"]}

        # argstep: i campi ereditati li decide il SERVER, non il client (a meno che la
        # slide non li dichiari editabili), e vengono congelati nel payload al momento
        # dell'invio. Da li in poi la risposta e autosufficiente.
        if slide["type"] == "argstep":
            cfg = json.loads(slide["config"])
            fields = _argstep_fields(cfg)
            collect = fields[-1]
            val = str(body.payload.get(collect) or "").strip()
            if not val:
                raise HTTPException(400, f"campo '{collect}' vuoto")
            carry = _resolve_carry(conn, rid, slide, body.token)
            editable = bool(cfg.get("carry_editable"))
            out = {collect: val}
            for f in fields[:-1]:
                sent = str(body.payload.get(f) or "").strip()
                inherited = carry["values"].get(f, "") if carry else ""
                out[f] = inherited if (inherited and not editable) else (sent or inherited)
                if not out[f]:
                    raise HTTPException(400, f"campo '{f}' mancante")
            src = carry["source_id"] if carry else None
            if cfg.get("single", True):
                conn.execute(
                    "DELETE FROM response WHERE run_id=? AND slide_id=? AND participant_token=?",
                    (rid, body.slide_id, body.token),
                )
            conn.execute(
                "INSERT INTO response (id, run_id, slide_id, participant_token, payload, "
                "source_response_id, created_at) VALUES (?,?,?,?,?,?,?)",
                (db.new_id(), rid, body.slide_id, body.token, json.dumps(out), src, db.now_iso()),
            )
            return {"ok": True, "single": bool(cfg.get("single", True))}

        # voto singolo → upsert: rimuovo il voto precedente di questo token.
        # `config.single` lo rende disponibile anche ai tipi testuali (opentext, argpoll).
        if slide["type"] in SINGLE_VOTE_TYPES or json.loads(slide["config"]).get("single"):
            conn.execute(
                "DELETE FROM response WHERE run_id=? AND slide_id=? AND participant_token=?",
                (rid, body.slide_id, body.token),
            )
        conn.execute(
            "INSERT INTO response (id, run_id, slide_id, participant_token, payload, created_at) "
            "VALUES (?,?,?,?,?,?)",
            (db.new_id(), rid, body.slide_id, body.token, json.dumps(body.payload), db.now_iso()),
        )
        return {"ok": True}


class AssignIn(BaseModel):
    slide_id: str
    token: str


@app.post("/api/live/{code}/assign")
def assign(code: str, body: AssignIn):
    """Randomizzatore: assegna il partecipante a un gruppo (bilanciato, stabile per token)."""
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] != "groups":
            raise HTTPException(404, "slide non valida")
        # assegnazione esistente → stabile
        ex = conn.execute(
            "SELECT payload FROM response WHERE run_id=? AND slide_id=? AND participant_token=?",
            (rid, body.slide_id, body.token),
        ).fetchone()
        if ex:
            return json.loads(ex["payload"])
        if _run_slide_state(conn, rid, body.slide_id) == "closed":
            raise HTTPException(409, "assegnazioni chiuse")
        groups = json.loads(slide["config"]).get("groups", [])
        if not groups:
            raise HTTPException(400, "nessun gruppo definito")
        # bilanciamento: assegna al gruppo attualmente meno numeroso
        counts = {g["id"]: 0 for g in groups}
        for r in conn.execute(
            "SELECT payload FROM response WHERE run_id=? AND slide_id=?", (rid, body.slide_id)
        ).fetchall():
            gid = json.loads(r["payload"]).get("group_id")
            if gid in counts:
                counts[gid] += 1
        chosen = min(groups, key=lambda g: counts[g["id"]])
        payload = {"group_id": chosen["id"], "group_name": chosen["name"]}
        conn.execute(
            "INSERT INTO response (id, run_id, slide_id, participant_token, payload, created_at) "
            "VALUES (?,?,?,?,?,?)",
            (db.new_id(), rid, body.slide_id, body.token, json.dumps(payload), db.now_iso()),
        )
        return payload


# ----------------------------------------------------------------------------
# Moonshot — energizer collaborativo (lifecycle + progresso on-demand)
# ----------------------------------------------------------------------------
def _moonshot_config(slide) -> dict:
    c = json.loads(slide["config"])
    return {
        "distance": max(1, int(c.get("distance", 8))),
        "reserve": max(0, int(c.get("reserve", 4))),
        "window_ms": int(c.get("window_ms", 3000)),
        "cycle_ms": int(c.get("cycle_ms", 5000)),
    }


def _ensure_moonshot_lobby(conn, run_id: str, slide_id: str) -> None:
    conn.execute(
        "INSERT OR IGNORE INTO moonshot_game (run_id, slide_id, status, total_crew) VALUES (?,?,'lobby',0)",
        (run_id, slide_id),
    )


def _moonshot_state(conn, run_id: str, slide, token: str | None = None) -> dict:
    """Stato completo del gioco, calcolato on-demand. `token` personalizza il ruolo."""
    sid = slide["id"]
    cfg = _moonshot_config(slide)
    N, X, win, cyc = cfg["distance"], cfg["reserve"], cfg["window_ms"], cfg["cycle_ms"]
    total_windows = N + X
    g = conn.execute(
        "SELECT * FROM moonshot_game WHERE run_id=? AND slide_id=?", (run_id, sid)
    ).fetchone()
    ready = conn.execute(
        "SELECT COUNT(*) AS c FROM moonshot_player WHERE run_id=? AND slide_id=?", (run_id, sid)
    ).fetchone()["c"]

    def _role():
        if not (g and token):
            return "spectator"
        if token == g["captain_token"]:
            return "captain"
        if conn.execute(
            "SELECT 1 FROM moonshot_player WHERE run_id=? AND slide_id=? AND token=?",
            (run_id, sid, token),
        ).fetchone():
            return "crew"
        return "spectator"

    out = {
        "config": cfg,
        "ready_count": ready,
        "role": _role(),
        "min_players": MOONSHOT_MIN_PLAYERS,
    }
    if not g or g["status"] == "lobby":
        out["status"] = "lobby"
        return out

    total_crew = g["total_crew"] or 1
    start = g["started_at_ms"]
    now = db.now_ms()
    elapsed = max(0, now - start)
    cur_idx = elapsed // cyc
    in_window = (elapsed % cyc) < win
    counts = {
        r["window_idx"]: r["c"]
        for r in conn.execute(
            "SELECT window_idx, COUNT(*) AS c FROM moonshot_boost "
            "WHERE run_id=? AND slide_id=? GROUP BY window_idx", (run_id, sid)
        ).fetchall()
    }
    progress = 0.0
    windows_used = 0
    last_ratio = 0.0
    for i in range(total_windows):
        if now >= start + i * cyc + win:            # finestra i chiusa
            ratio = min(1.0, counts.get(i, 0) / total_crew)
            progress += ratio
            windows_used += 1
            last_ratio = ratio
        else:
            break
    won = progress >= N
    failed = windows_used >= total_windows and not won
    result = "success" if won else ("failed" if failed else None)
    frac = min(1.0, progress / N)
    open_idx = int(cur_idx) if (in_window and cur_idx < total_windows and result is None) else None
    you_boosted = bool(
        token and open_idx is not None and conn.execute(
            "SELECT 1 FROM moonshot_boost WHERE run_id=? AND slide_id=? AND window_idx=? AND token=?",
            (run_id, sid, open_idx, token),
        ).fetchone()
    )
    out.update({
        "status": "running",
        "total_crew": total_crew,
        "captain": g["captain_token"],
        "progress": round(progress, 4),
        "progress_frac": round(frac, 4),
        "altitude_km": round(frac * MOON_DISTANCE_KM),
        "boosts_left": max(0, total_windows - windows_used),
        "windows_used": windows_used,
        "total_windows": total_windows,
        "room_energy": round(last_ratio, 3),
        "result": result,
        "window_open": open_idx is not None,
        "window_idx": open_idx,
        "you_boosted": you_boosted,
        # per il countdown locale del capitano (schedule ricostruibile client-side)
        "started_at_ms": start,
        "server_now_ms": now,
    })
    return out


class MoonshotIn(BaseModel):
    slide_id: str
    token: str


class MoonshotSlideIn(BaseModel):
    slide_id: str


@app.post("/api/live/{code}/moonshot/ready")
def moonshot_ready(code: str, body: MoonshotIn):
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        _ensure_moonshot_lobby(conn, rid, body.slide_id)
        g = conn.execute(
            "SELECT status FROM moonshot_game WHERE run_id=? AND slide_id=?", (rid, body.slide_id)
        ).fetchone()
        if g["status"] != "lobby":
            return {"ok": True, "ready": False, "spectator": True}  # partita avviata → spettatore
        conn.execute(
            "INSERT OR IGNORE INTO moonshot_player (run_id, slide_id, token) VALUES (?,?,?)",
            (rid, body.slide_id, body.token),
        )
        return {"ok": True, "ready": True}


@app.post("/api/live/{code}/moonshot/boost")
def moonshot_boost(code: str, body: MoonshotIn):
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            raise HTTPException(409, "nessun run attivo")
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] != "moonshot":
            raise HTTPException(404, "slide non valida")
        st = _moonshot_state(conn, rid, slide, body.token)
        if st["status"] != "running" or st.get("result") is not None:
            raise HTTPException(409, "gioco non attivo")
        if not st.get("window_open"):
            return {"ok": True, "boosted": False}   # nessuna finestra aperta: tap ignorato (soft)
        if not conn.execute(
            "SELECT 1 FROM moonshot_player WHERE run_id=? AND slide_id=? AND token=?",
            (rid, body.slide_id, body.token),
        ).fetchone():
            raise HTTPException(403, "non sei nella crew")
        conn.execute(
            "INSERT OR IGNORE INTO moonshot_boost (run_id, slide_id, window_idx, token) VALUES (?,?,?,?)",
            (rid, body.slide_id, st["window_idx"], body.token),
        )
        return {"ok": True, "boosted": True}


@app.post("/api/runs/{rid}/moonshot/launch")
def moonshot_launch(rid: str, body: MoonshotSlideIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        _ensure_moonshot_lobby(conn, rid, body.slide_id)
        players = [
            r["token"] for r in conn.execute(
                "SELECT token FROM moonshot_player WHERE run_id=? AND slide_id=?", (rid, body.slide_id)
            ).fetchall()
        ]
        if len(players) < MOONSHOT_MIN_PLAYERS:
            raise HTTPException(400, f"servono almeno {MOONSHOT_MIN_PLAYERS} giocatori pronti")
        conn.execute(
            "UPDATE moonshot_game SET status='running', captain_token=?, total_crew=?, started_at_ms=? "
            "WHERE run_id=? AND slide_id=?",
            (secrets.choice(players), len(players), db.now_ms(), rid, body.slide_id),
        )
        return {"ok": True, "total_crew": len(players)}


@app.post("/api/runs/{rid}/moonshot/reassign")
def moonshot_reassign(rid: str, body: MoonshotSlideIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        players = [
            r["token"] for r in conn.execute(
                "SELECT token FROM moonshot_player WHERE run_id=? AND slide_id=?", (rid, body.slide_id)
            ).fetchall()
        ]
        if not players:
            raise HTTPException(400, "nessun giocatore")
        conn.execute(
            "UPDATE moonshot_game SET captain_token=? WHERE run_id=? AND slide_id=?",
            (secrets.choice(players), rid, body.slide_id),
        )
        return {"ok": True}


@app.post("/api/runs/{rid}/moonshot/reset")
def moonshot_reset(rid: str, body: MoonshotSlideIn, user: dict = CurrentUser):
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        for tbl in ("moonshot_boost", "moonshot_player", "moonshot_game"):
            conn.execute(f"DELETE FROM {tbl} WHERE run_id=? AND slide_id=?", (rid, body.slide_id))
        return {"ok": True}


class ClusterIn(BaseModel):
    slide_id: str


def _record_usage(conn, user_id: str, run_id: str, slide_id: str, kind: str, usage: dict) -> None:
    """Registra il consumo della chiave centrale in usage_log, con costo stimato (prezzi Sonnet)."""
    it = int(usage.get("input", 0))
    ot = int(usage.get("output", 0))
    cost = it / 1_000_000 * PRICE_IN + ot / 1_000_000 * PRICE_OUT
    conn.execute(
        "INSERT INTO usage_log (id, user_id, run_id, slide_id, kind, model, "
        "input_tokens, output_tokens, cost_usd, created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
        (db.new_id(), user_id, run_id, slide_id, kind, clustering.MODEL, it, ot, cost, db.now_iso()),
    )


# --- timer: una slide che dura, invece di una che chiede ---------------------
# L'unico tipo il cui contenuto e' il tempo. Vive con lo stesso patto di moonshot:
# il server tiene l'istante, il client conta da solo. La differenza e' che qui si
# salva la FINE e non l'inizio - cosi' aggiungere due minuti (che in aula capita
# sempre) e' un UPDATE di una colonna, e non un ricalcolo di durata.

TIMER_DEFAULT_MIN = 15.0
TIMER_MAX_MIN = 600.0


def _timer_config(slide) -> dict:
    raw = slide["config"]
    cfg = json.loads(raw) if isinstance(raw, str) else (raw or {})
    try:
        minutes = float(cfg.get("minutes", TIMER_DEFAULT_MIN))
    except (TypeError, ValueError):
        minutes = TIMER_DEFAULT_MIN
    minutes = max(0.1, min(TIMER_MAX_MIN, minutes))
    return {
        "minutes": minutes,
        "autostart": cfg.get("autostart", True) is not False,
        "show_end_time": cfg.get("show_end_time", True) is not False,
    }


def _timer_start(conn, run_id: str, slide, only_if_absent: bool = False) -> None:
    """(Ri)avvia il countdown adesso. Con `only_if_absent` non tocca uno gia' in corso."""
    cfg = _timer_config(slide)
    now = db.now_ms()
    ends = now + int(cfg["minutes"] * 60_000)
    if only_if_absent:
        # tornare sulla slide di pausa non fa ripartire il tempo: chi e' fuori dalla
        # stanza sta guardando quel numero, e vederlo risalire e' peggio che non averlo
        conn.execute(
            "INSERT INTO timer_state (run_id, slide_id, started_at_ms, ends_at_ms) "
            "VALUES (?,?,?,?) ON CONFLICT(run_id, slide_id) DO NOTHING",
            (run_id, slide["id"], now, ends),
        )
    else:
        conn.execute(
            "INSERT INTO timer_state (run_id, slide_id, started_at_ms, ends_at_ms) "
            "VALUES (?,?,?,?) ON CONFLICT(run_id, slide_id) DO UPDATE SET "
            "started_at_ms=excluded.started_at_ms, ends_at_ms=excluded.ends_at_ms",
            (run_id, slide["id"], now, ends),
        )


def _timer_state(conn, run_id: str, slide) -> dict:
    """Stato del countdown. `server_now_ms` viaggia sempre: e' cio' che permette al
    client di correggere lo scarto fra il proprio orologio e quello del server."""
    cfg = _timer_config(slide)
    row = conn.execute(
        "SELECT started_at_ms, ends_at_ms FROM timer_state WHERE run_id=? AND slide_id=?",
        (run_id, slide["id"]),
    ).fetchone()
    now = db.now_ms()
    out = {"config": cfg, "server_now_ms": now}
    if not row:
        out["status"] = "idle"
        return out
    out.update({
        "status": "running",
        "started_at_ms": row["started_at_ms"],
        "ends_at_ms": row["ends_at_ms"],
        "remaining_ms": max(0, row["ends_at_ms"] - now),
        "expired": now >= row["ends_at_ms"],
    })
    return out


class TimerIn(BaseModel):
    slide_id: str
    seconds: int | None = None      # solo per `extend`


@app.post("/api/runs/{rid}/timer/{action}")
def timer_action(rid: str, action: str, body: TimerIn, user: dict = CurrentUser):
    """start = (ri)avvia da adesso, extend = sposta la fine, clear = torna a non avviato."""
    if action not in ("start", "extend", "clear"):
        raise HTTPException(400, "azione non valida")
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] != "timer":
            raise HTTPException(404, "slide timer non trovata")
        if action == "start":
            _timer_start(conn, rid, slide)
        elif action == "clear":
            conn.execute(
                "DELETE FROM timer_state WHERE run_id=? AND slide_id=?", (rid, body.slide_id)
            )
        else:
            secs = int(body.seconds or 0)
            if not secs:
                raise HTTPException(400, "extend richiede `seconds`")
            row = conn.execute(
                "SELECT ends_at_ms FROM timer_state WHERE run_id=? AND slide_id=?",
                (rid, body.slide_id),
            ).fetchone()
            if not row:
                raise HTTPException(409, "timer non avviato")
            # si estende dalla fine se il tempo c'e' ancora, da adesso se e' gia' scaduto:
            # "+2 minuti" detto a tempo scaduto vuol dire due minuti da ora, non due minuti fa
            base = max(row["ends_at_ms"], db.now_ms())
            conn.execute(
                "UPDATE timer_state SET ends_at_ms=? WHERE run_id=? AND slide_id=?",
                (base + secs * 1000, rid, body.slide_id),
            )
        return _timer_state(conn, rid, slide)


@app.post("/api/runs/{rid}/cluster")
def cluster_run_slide(rid: str, body: ClusterIn, user: dict = CurrentUser):
    """Clusterizza (LLM) le risposte argpoll/opentext del run.
    Tier free → chiave dell'utente; tier full/admin → chiave centrale del server (con cost tracking)."""
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        role = user.get("role", "free")
        central = role in ("full", "admin")
        if central:
            key = CENTRAL_API_KEY
            if not key:
                raise HTTPException(503, "chiave API centrale non configurata sul server")
        else:  # free: chiave propria
            row = conn.execute("SELECT api_key FROM user WHERE id=?", (user["id"],)).fetchone()
            key = row["api_key"] if row else None
            if not key:
                raise HTTPException(400, "API key non configurata")
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] not in ("argpoll", "opentext", "argstep"):
            raise HTTPException(404, "slide non valida")
        rows = conn.execute(
            "SELECT payload FROM response WHERE run_id=? AND slide_id=? AND status='visible' "
            "ORDER BY created_at",
            (rid, body.slide_id),
        ).fetchall()
        if len(rows) < 2:
            raise HTTPException(400, "servono almeno 2 risposte per clusterizzare")
        question = slide["question"]
        try:
            if slide["type"] == "argstep":
                # la tappa finale porta tutti e tre i campi: gli assi sono claim x obiezione.
                # Le tappe intermedie ricadono sui prompt gia esistenti.
                fields = _argstep_fields(json.loads(slide["config"]))
                payloads = [json.loads(r["payload"]) for r in rows]
                if len(fields) >= 3:
                    recs = [{"n": i, "claim": p.get("claim", ""),
                             "justification": p.get("justification", ""),
                             "objection": p.get("objection", "")}
                            for i, p in enumerate(payloads, start=1)]
                    result, usage = clustering.cluster_argstep_full(key, question, recs)
                elif len(fields) == 2:
                    recs = [{"n": i, "claim": p.get("claim", ""),
                             "justification": p.get("justification", "")}
                            for i, p in enumerate(payloads, start=1)]
                    result, usage = clustering.cluster_argpoll(key, question, recs)
                else:
                    recs = [{"n": i, "text": p.get("claim", "")}
                            for i, p in enumerate(payloads, start=1)]
                    result, usage = clustering.cluster_argstep_claims(key, question, recs)
            elif slide["type"] == "argpoll":
                pairs = []
                for i, r in enumerate(rows, start=1):
                    p = json.loads(r["payload"])
                    pairs.append({"n": i, "claim": p.get("claim", ""), "justification": p.get("justification", "")})
                result, usage = clustering.cluster_argpoll(key, question, pairs)
            else:  # opentext
                texts = []
                for i, r in enumerate(rows, start=1):
                    texts.append({"n": i, "text": json.loads(r["payload"]).get("text", "")})
                result, usage = clustering.cluster_opentext(key, question, texts)
        except Exception as e:  # errore LLM / chiave / parsing
            raise HTTPException(502, f"clustering fallito: {e}")
        if central:  # traccia il consumo della chiave centrale (cost tracking per-utente)
            _record_usage(conn, user["id"], rid, body.slide_id, slide["type"], usage)
        if slide["type"] == "argstep":
            _materialize_clusters(conn, rid, body.slide_id, result)
            return _argstep_results(conn, rid, slide)
        if slide["type"] == "argpoll":
            _materialize_clusters(conn, rid, body.slide_id, result)
            return _argpoll_clustered(conn, rid, body.slide_id)
        _materialize_text_clusters(conn, rid, body.slide_id, result)
        return _opentext_clustered(conn, rid, body.slide_id)


class ArgstepAddIn(BaseModel):
    slide_id: str
    source_response_id: str       # a quale riga della tabella si attacca l'obiezione
    text: str


@app.post("/api/runs/{rid}/argstep/add")
def argstep_add(rid: str, body: ArgstepAddIn, user: dict = CurrentUser):
    """Il docente trascrive un'obiezione arrivata a voce dall'aula. Entra nella stessa
    tabella e nello stesso clustering delle altre, marcata `by: presenter`.
    Token usa e getta: non deve sovrascrivere le precedenti aggiunte del docente."""
    text = body.text.strip()
    if not text:
        raise HTTPException(400, "obiezione vuota")
    with db.get_conn() as conn:
        _check_owner(conn, _pid_of_run(conn, rid), user)
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (body.slide_id,)).fetchone()
        if not slide or slide["type"] != "argstep":
            raise HTTPException(404, "slide non valida")
        fields = _argstep_fields(json.loads(slide["config"]))
        if fields[-1] != "objection":
            raise HTTPException(400, "questa tappa non raccoglie obiezioni")
        src = conn.execute(
            "SELECT id, payload FROM response WHERE id=? AND run_id=?",
            (body.source_response_id, rid),
        ).fetchone()
        if not src:
            raise HTTPException(404, "riga sorgente non trovata")
        sp = json.loads(src["payload"])
        payload = {
            "claim": sp.get("claim", ""),
            "justification": sp.get("justification", ""),
            "objection": text,
            "by": "presenter",
        }
        conn.execute(
            "INSERT INTO response (id, run_id, slide_id, participant_token, payload, "
            "source_response_id, created_at) VALUES (?,?,?,?,?,?,?)",
            (db.new_id(), rid, body.slide_id, "presenter:" + db.new_id(),
             json.dumps(payload), src["id"], db.now_iso()),
        )
        return {"ok": True}


@app.get("/api/live/{code}/results")
def live_results(code: str):
    with db.get_conn() as conn:
        pres = _resolve_run(conn, code)
        rid = pres["active_run_id"]
        if not rid:
            return JSONResponse({"status": "waiting"})
        run = conn.execute("SELECT * FROM run WHERE id=?", (rid,)).fetchone()
        active = run["active_slide_id"]
        if not active:
            return JSONResponse({"status": "waiting"})
        slide = conn.execute("SELECT * FROM slide WHERE id=?", (active,)).fetchone()
        if slide is None:
            return JSONResponse({"status": "waiting"})
        res = _results(conn, rid, slide)
        # quiz: non svelare la risposta corretta finché la slide non è rivelata
        if res.get("quiz") and _run_slide_state(conn, rid, active) != "revealed":
            res = {**res, "correct": []}
        return res
